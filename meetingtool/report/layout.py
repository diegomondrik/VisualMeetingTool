"""What a company's template holds besides its design (INGOL D-188): the
cover's fields, where the report starts, and its table of contents.

A field is a name in braces the company writes where the meeting's data goes
({cliente}, {fecha}...), in the body or in a header or footer. Word often
splits what one types into several runs (spelling, an edit, a change of
format), so a field is found in a paragraph's text and replaced across the
runs it spans, keeping the first run's format.

What a template has after its table of contents, or after {informe} alone on
a line, is a model of a report, not part of every report: it is dropped.

A table of contents is a Word TOC field. What Word saved in it is replaced by
the headings of the report, each a link to a bookmark on its heading. There
are no page numbers: they need the document laid out, which Word does when
the reader updates the field. The document is not marked to update its fields
on opening, which would make Word ask every reader for permission.
"""

import dataclasses
import datetime
import io
import re

import docx
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

FIELDS = {"cliente": "client", "client": "client", "proyecto": "project", "project": "project",
          "reunion": "meeting", "reunión": "meeting", "meeting": "meeting", "fecha": "date", "date": "date",
          "tipo": "type", "type": "type"}
REPORT_START = frozenset({"informe", "report"})
FIELD = re.compile(r"\{\s*([^\W\d_]+)\s*\}")
KNOWN = "{cliente} {proyecto} {reunion} {fecha} {tipo} (or {client} {project} {meeting} {date} {type})"

MONTHS = {"es": ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
                 "octubre", "noviembre", "diciembre"],
          "en": ["January", "February", "March", "April", "May", "June", "July", "August", "September",
                 "October", "November", "December"]}
TYPES = {"es": {"presale": "Preventa", "negotiation": "Venta o negociación", "requirements": "Relevamiento",
                "kickoff": "Inicio de proyecto", "status": "Seguimiento", "technical": "Técnica",
                "training": "Capacitación", "discovery": "Descubrimiento"},
         "en": {"presale": "Presale", "negotiation": "Sales or negotiation", "requirements": "Requirements",
                "kickoff": "Kickoff", "status": "Status", "technical": "Technical", "training": "Training",
                "discovery": "Discovery"}}
LEVELS = re.compile(r'\\o\s+"(\d)-(\d)"')
BOOKMARK_PREFIX = "_TocMeetingTool"
INDENT = Cm(0.63)


class LayoutError(Exception):
    """A template whose fields or report start cannot be used."""


@dataclasses.dataclass
class Layout:
    start: int | None      # index of the body's first dropped child; None drops nothing
    start_kind: str | None  # "index", "marker" or None
    tocs: list              # (begin paragraph, end paragraph, instruction) of each kept table of contents
    fields: list            # kinds of the fields found in what is kept, in order of first appearance
    dropped: int            # paragraphs with text among the dropped children


# ── Reading the template ─────────────────────────────────────────────────────

def body_children(document):
    return [child for child in document.element.body.iterchildren() if child.tag != qn("w:sectPr")]


def _paragraph_of(element):
    while element is not None and element.tag != qn("w:p"):
        element = element.getparent()
    return element


def _top(element, body):
    while element.getparent() is not body:
        element = element.getparent()
    return element


def _tocs(document):
    """Each table of contents in the body: (begin paragraph, end paragraph,
    instruction). A TOC field that lists figures or tables (\\c, \\f) is not
    one."""
    found, stack = [], []
    for node in document.element.body.iter(qn("w:fldChar"), qn("w:instrText")):
        if node.tag == qn("w:instrText"):
            if stack:
                stack[-1]["instr"] += node.text or ""
            continue
        kind = node.get(qn("w:fldCharType"))
        if kind == "begin":
            stack.append({"begin": _paragraph_of(node), "instr": ""})
        elif kind == "end" and stack:
            field = stack.pop()
            instr = field["instr"].strip()
            words = instr.split()
            if not stack and words and words[0].upper() == "TOC" and not re.search(r"\\[cf]\b", instr):
                found.append((field["begin"], _paragraph_of(node), instr))
    return found


def paragraphs(elements):
    """Every paragraph in the elements, tables and text boxes included."""
    for element in elements:
        if element.tag == qn("w:p"):
            yield element
        yield from (p for p in element.iter(qn("w:p")) if p is not element)


def header_footer_parts(document):
    seen = set()
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if part.is_linked_to_previous or id(part._element) in seen:
                continue
            seen.add(id(part._element))
            yield part._element


def _texts(element):
    return [node for node in element.iter(qn("w:t")) if _paragraph_of(node) is element]


def paragraph_text(element):
    return "".join(node.text or "" for node in _texts(element))


def read_layout(document):
    """Where the report starts in the template, what it keeps, and the
    fields in what is kept; LayoutError for a field it does not know or a
    report start that is not alone on a line of the body."""
    body = document.element.body
    children = body_children(document)
    index = {id(child): number for number, child in enumerate(children)}
    start, start_kind = None, None
    for number, child in enumerate(children):
        if child.tag == qn("w:p"):
            names = [match.group(1).lower() for match in FIELD.finditer(paragraph_text(child))]
            if any(name in REPORT_START for name in names):
                if FIELD.fullmatch(paragraph_text(child).strip()) is None:
                    raise LayoutError("{informe} must be alone on its line, with nothing else on it")
                start, start_kind = number, "marker"
                break
    tocs = _tocs(document)
    if start is None and tocs:
        start, start_kind = index[id(_top(tocs[-1][1], body))] + 1, "index"
    kept = children if start is None else children[:start]
    dropped = children[len(kept):]
    kept_ids = {id(child) for child in kept}
    tocs = [toc for toc in tocs if id(_top(toc[0], body)) in kept_ids]
    for begin, end, _ in tocs:
        separate = [char for char in begin.iter(qn("w:fldChar")) if char.get(qn("w:fldCharType")) == "separate"]
        if begin.getparent() is not end.getparent() or not separate:
            raise LayoutError("the table of contents cannot be read: insert it again in Word "
                              "(References, Table of Contents) and save the template")
    fields, unknown, misplaced = [], [], False
    for element in list(paragraphs(kept)) + [p for part in header_footer_parts(document) for p in paragraphs([part])]:
        for match in FIELD.finditer(paragraph_text(element)):
            name = match.group(1).lower()
            if name in REPORT_START:
                misplaced = True  # one alone on a line of the body starts the report and is not kept
            elif name in FIELDS:
                if FIELDS[name] not in fields:
                    fields.append(FIELDS[name])
            elif match.group(0) not in unknown:
                unknown.append(match.group(0))
    if misplaced:
        raise LayoutError("{informe} must be alone on its line in the body, not in a table, header or footer")
    if unknown:
        raise LayoutError(f"the template has {', '.join(unknown)}, which is not a field it knows; the fields are "
                          f"{KNOWN}")
    model = dropped[1:] if start_kind == "marker" else dropped  # the marker's own line is not model
    dropped_text = sum(1 for p in paragraphs(model) if paragraph_text(p).strip())
    return Layout(start, start_kind, tocs, fields, dropped_text)


# ── Filling it ───────────────────────────────────────────────────────────────

def drop_model(document, layout):
    """Remove the template's model; return the section break that ended the
    cover, when the first one is among what was removed (the cover is then a
    section of its own), else None."""
    if layout.start is None:
        return None
    body, kept = document.element.body, None
    for child in body_children(document)[layout.start:]:
        properties = child.find(qn("w:pPr")) if child.tag == qn("w:p") else None
        section = properties.find(qn("w:sectPr")) if properties is not None else None
        if section is not None and kept is None and not _section_before(document, layout.start):
            properties.remove(section)
            kept = section
        body.remove(child)
    return kept


def _section_before(document, start):
    """Whether a section break already ends somewhere in the kept body."""
    return any(child.tag == qn("w:p") and child.find(qn("w:pPr")) is not None
               and child.find(qn("w:pPr")).find(qn("w:sectPr")) is not None
               for child in body_children(document)[:start])


def field_values(language, *, client=None, project=None, meeting=None, date=None, meeting_type=None):
    written = ""
    if date:
        day = datetime.date.fromisoformat(date)
        month = MONTHS[language][day.month - 1]
        written = f"{day.day} de {month} de {day.year}" if language == "es" else f"{month} {day.day}, {day.year}"
    return {"client": client or "", "project": project or "", "meeting": meeting or "", "date": written,
            "type": TYPES[language].get(meeting_type or "", meeting_type or "")}


def _fill_paragraph(element, values):
    nodes = _texts(element)
    text = "".join(node.text or "" for node in nodes)
    matches = [m for m in FIELD.finditer(text) if m.group(1).lower() in FIELDS]
    for match in reversed(matches):
        value = values[FIELDS[match.group(1).lower()]]
        offset, spans = 0, []
        for node in nodes:
            length = len(node.text or "")
            if offset < match.end() and offset + length > match.start():
                spans.append((node, max(match.start() - offset, 0), min(match.end() - offset, length)))
            offset += length
        first = spans[0][0]
        for node, begin, end in spans:
            node.text = (node.text[:begin] + (value if node is first else "") + node.text[end:])
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def fill_fields(document, values):
    for element in paragraphs(body_children(document)):
        _fill_paragraph(element, values)
    for part in header_footer_parts(document):
        for element in paragraphs([part]):
            _fill_paragraph(element, values)


def leftover_fields(document, cover_children):
    """The known fields still written in the cover or a header or footer."""
    elements = list(paragraphs(cover_children))
    elements += [p for part in header_footer_parts(document) for p in paragraphs([part])]
    return sorted({match.group(0) for p in elements for match in FIELD.finditer(paragraph_text(p))
                   if match.group(1).lower() in FIELDS})


def no_update_on_open(document):
    """Word would ask the reader for permission to update fields."""
    settings = document.settings.element
    for node in settings.findall(qn("w:updateFields")):
        settings.remove(node)


# ── The table of contents ────────────────────────────────────────────────────

def levels(instruction):
    match = LEVELS.search(instruction)
    return (int(match.group(1)), int(match.group(2))) if match else (1, 9)


def _toc_style(document, level):
    for style in document.styles.element.iterchildren(qn("w:style")):
        name = style.find(qn("w:name"))
        if name is not None and (name.get(qn("w:val")) or "").lower() == f"toc {level}":
            return style.get(qn("w:styleId"))
    return None


def _bookmark_ids(document):
    return [int(node.get(qn("w:id"))) for node in document.element.body.iter(qn("w:bookmarkStart"))
            if (node.get(qn("w:id")) or "").isdigit()]


def bookmark_headings(document, headings):
    """Put a bookmark around each heading paragraph; return their names."""
    next_id = max(_bookmark_ids(document), default=0) + 1
    names = []
    for number, (_, _, element) in enumerate(headings, start=1):
        name = f"{BOOKMARK_PREFIX}{number:04d}"
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(next_id))
        start.set(qn("w:name"), name)
        end.set(qn("w:id"), str(next_id))
        properties = element.find(qn("w:pPr"))
        position = 1 if properties is not None else 0
        element.insert(position, start)
        element.append(end)
        names.append(name)
        next_id += 1
    return names


def _run(child):
    run = OxmlElement("w:r")
    run.append(child)
    return run


def _field_char(kind):
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), kind)
    return node


def _entry(document, level, text, bookmark):
    paragraph = OxmlElement("w:p")
    properties = paragraph.get_or_add_pPr()
    style = _toc_style(document, level)
    if style:
        properties.style = style
    elif level > 1:
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), str(int(INDENT.twips * (level - 1))))
        properties.append(indent)
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), bookmark)
    link.set(qn("w:history"), "1")
    node = OxmlElement("w:t")
    node.text = text
    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    link.append(_run(node))
    paragraph.append(link)
    return paragraph


def fill_toc(document, toc, headings, bookmarks):
    """Replace what Word saved in the table of contents with the report's
    headings within its levels; return the entries' texts."""
    begin, end, instruction = toc
    low, high = levels(instruction)
    entries = [(level, text, name) for (level, text, _), name in zip(headings, bookmarks) if low <= level <= high]
    opening, before = [], []
    state = "before"
    for run in begin.iterchildren():
        if run.tag == qn("w:pPr"):
            continue
        chars = run.findall(qn("w:fldChar")) if run.tag == qn("w:r") else []
        kinds = [c.get(qn("w:fldCharType")) for c in chars]
        if state == "before" and "begin" in kinds:
            state = "field"
        if state == "before":
            before.append(run)
        elif state == "field":
            opening.append(run)
            if "separate" in kinds:
                state = "result"
                break
    if state != "result":
        raise LayoutError("the table of contents cannot be read: its field has no result part")
    for char in opening[0].findall(qn("w:fldChar")):
        char.attrib.pop(qn("w:dirty"), None)
    after = []
    closing = [child for child in end.iterchildren() if child.tag == qn("w:r")
               and any(c.get(qn("w:fldCharType")) == "end" for c in child.findall(qn("w:fldChar")))]
    if closing:
        found = False
        for child in end.iterchildren():
            if found and child.tag != qn("w:pPr"):
                after.append(child)
            if child is closing[-1]:
                found = True
    parent = begin.getparent()
    position = parent.index(begin)
    siblings = list(parent.iterchildren())
    for child in siblings[position:siblings.index(end) + 1]:
        parent.remove(child)
    made = [_entry(document, level, text, name) for level, text, name in entries] or [OxmlElement("w:p")]
    first = made[0]
    anchor = 1 if first.find(qn("w:pPr")) is not None else 0
    for offset, node in enumerate(before + opening):
        first.insert(anchor + offset, node)
    made[-1].append(_run(_field_char("end")))
    for node in after:
        made[-1].append(node)
    properties = end.find(qn("w:pPr"))
    section = properties.find(qn("w:sectPr")) if properties is not None else None
    if section is not None:  # the table of contents ends a section: its last entry still does
        made[-1].get_or_add_pPr().append(section)
    for offset, paragraph in enumerate(made):
        parent.insert(position + offset, paragraph)
    return [text for _, text, _ in entries]


def toc_entries(document, toc):
    """The texts of a table of contents' entries: its links, one per
    paragraph (text before the field in its first paragraph is not an
    entry)."""
    begin, end, _ = toc
    parent = begin.getparent()
    siblings = list(parent.iterchildren())
    texts = []
    for element in siblings[siblings.index(begin):siblings.index(end) + 1]:
        if element.tag == qn("w:p"):
            texts.append("".join(node.text or "" for link in element.iter(qn("w:hyperlink"))
                                 for node in link.iter(qn("w:t"))).strip())
    return [text for text in texts if text]


# ── The example ──────────────────────────────────────────────────────────────

def example_template():
    """A template to start from: the five fields on the cover, a table of
    contents, and after it a note that is dropped from every report."""
    template = docx.Document()
    template.sections[0].header.paragraphs[0].text = "Nombre de su empresa"
    template.sections[0].footer.paragraphs[0].text = "www.su-empresa.example"
    template.add_paragraph("{proyecto}", style="Title")
    for label, field in (("Cliente", "{cliente}"), ("Reunión", "{reunion}"), ("Tipo de reunión", "{tipo}"),
                         ("Fecha", "{fecha}")):
        paragraph = template.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(field)
    template.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    template.add_paragraph("Índice", style="TOC Heading")
    field = template.add_paragraph()._p
    instruction = OxmlElement("w:instrText")
    instruction.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    placeholder = OxmlElement("w:t")
    placeholder.text = "(cada informe pone acá sus secciones)"
    for node in (_field_char("begin"), instruction, _field_char("separate"), placeholder, _field_char("end")):
        field.append(_run(node))
    template.add_paragraph("Todo lo que esté después del índice es un modelo y no entra en los informes.")
    template.add_paragraph(
        "Cómo usar esta plantilla: cambie el encabezado, el pie, los colores y las letras como quiera. En la "
        "portada, escriba entre llaves el dato que va en cada lugar: {cliente}, {proyecto}, {reunion}, {fecha} "
        "o {tipo}. El índice se llena con las secciones de cada informe; para ver los números de página, en "
        "Word: clic derecho sobre el índice, Actualizar campos.")
    buffer = io.BytesIO()
    template.save(buffer)
    return buffer.getvalue()
