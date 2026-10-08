"""Build the meeting report as a Word document (INGOL D-177).

The original MeetingTool exported its report to Word with python-docx
(tools/exporter.py): only the frames the report names are embedded, where
they are named, and a named frame that is missing stops the export. That is
kept. What is new: no network and no key, so the report can be rebuilt at no
cost after the summary is edited by hand; the document is opened again and
counted before it is delivered (the original left a note in place of an image
it could not embed and delivered anyway); a mention that looks like a frame
but names none stops it too; and the company that runs the analysis can give
its own Word template, one per installation, whose page becomes the cover and
whose header, footer and styles the report keeps.

A template saved by Word (measured on Word in Spanish) defines the heading
and title styles but no list or table style and no numbering, so lists and
tables here use direct formatting, never a named style a template may lack.
A template with macros is refused: the report is sent to third parties.
"""

import collections
import dataclasses
import hashlib
import io
import json
import os
import re
import time
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import docx
from docx.enum.text import WD_BREAK
from docx.image.exceptions import UnrecognizedImageError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from meetingtool import disk, texts, word_package
from meetingtool.projects import store
from meetingtool.report import layout
from meetingtool.summary import qa, writer

SUMMARY_NAME = writer.OUTPUT_NAME
OUTPUT_NAME = "summary.docx"
TEMPLATE_NAME = "report-template.docx"
TEMPLATE_RECORD = "report-template.json"
TEMPLATE_EXTENSIONS = {".docx", ".dotx"}
MACRO_EXTENSIONS = {".docm", ".dotm"}
DOCUMENT_TYPE = b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
TEMPLATE_TYPE = b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"

# Relationship types that bring an embedded object, a control or a macro
# project into the document, or another template or document into it.
ACTIVE_RELATIONSHIPS = frozenset({"attachedTemplate", "oleObject", "package", "control", "activeXControl",
                                  "activeXControlBinary", "vbaProject", "wordVbaData", "aFChunk", "subDocument",
                                  "frame"})
_ACTIVE_RELATIONSHIPS = {kind.lower() for kind in ACTIVE_RELATIONSHIPS}
# The only fields a template may hold (WI23): the page, the section, the table of contents and the references
# inside the document, the date and the document's properties, a condition, a sequence and the formula (=). Every
# other field is refused, so a field that brings outside content and that no list knew of cannot pass (WI22-P3-5
# of the limitations register). Compared in upper case.
ALLOWED_FIELDS = frozenset({"PAGE", "NUMPAGES", "SECTIONPAGES", "SECTION", "TOC", "PAGEREF", "REF", "NOTEREF",
                            "STYLEREF", "HYPERLINK", "DATE", "TIME", "CREATEDATE", "SAVEDATE", "PRINTDATE",
                            "DOCPROPERTY", "TITLE", "SUBJECT", "AUTHOR", "IF", "SEQ", "="})
# The longest name of a refused field that a message says: what comes from the file is not given whole.
NAME_SHOWN = 40
# Where a hyperlink may go: a web page, a mail address, or a place in the same document.
HYPERLINK_ALLOWED = re.compile(r"\s*(?:(?:https?|mailto):|#)", re.IGNORECASE)
# The switches of a HYPERLINK field that take the next word as their argument.
SWITCHES_WITH_ARGUMENT = frozenset({"\\l", "\\o", "\\t"})
# WordprocessingML's namespace as Word writes it (Transitional) and as ISO 29500 Strict names it.
WORD_NAMESPACES = frozenset({"http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                             "http://purl.oclc.org/ooxml/wordprocessingml/main"})
# The tags (as ElementTree writes them: {namespace}name) that carry a field, whatever prefix a part gives them.
FIELD_TAGS = {f"{{{namespace}}}{local}": local for namespace in WORD_NAMESPACES
              for local in ("fldSimple", "instrText", "delInstrText", "fldChar")}
# Markup Compatibility (ISO 29500-3): how a producer offers Word alternatives (mc:AlternateContent, whose
# mc:Choice and mc:Fallback Word reads one of) and names the namespaces Word may skip, element and content
# (mc:Ignorable). Which it reads depends on what that Word understands, which the file does not say.
MC_NAMESPACE = "http://schemas.openxmlformats.org/markup-compatibility/2006"
MC_ELEMENTS = frozenset({"AlternateContent", "Choice", "Fallback"})
MC_PREFIX_ATTRIBUTES = frozenset({"Ignorable", "MustUnderstand"})
MC_NAME_ATTRIBUTES = frozenset({"ProcessContent", "PreserveElements", "PreserveAttributes"})  # prefix:name items
# Parts that hold XML even when their name does not say so are found by content type; these are the
# extensions that always do.
XML_EXTENSIONS = (".xml", ".rels")
# What a content type holds when the package carries macros (compared in lower case).
MACRO_TYPES = ("macroenabled", "vbaproject", "vbadata")

FRAME_REF = re.compile(r"\[(frame_\d+_t(\d{2})-(\d{2})-(\d{2})\.jpg)\]")
FRAME_LIKE = re.compile(r"\bframes?_", re.IGNORECASE)
# As the summary's own check reads a heading (meetingtool.summary.writer):
# up to the text, the space after the hashes is optional.
HEADING = re.compile(r"^(#{1,6})\s*(\S.*?)\s*#*\s*$")
BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBERED = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
TABLE_SEPARATOR = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")
INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|`[^`]+`)")

LABELS = {
    "es": {"title": "Resumen de la reunión", "date": "Fecha", "project": "Proyecto",
           "mention": "imagen {clock}", "caption": "Imagen del minuto {clock} de la reunión"},
    "en": {"title": "Meeting summary", "date": "Date", "project": "Project",
           "mention": "frame {clock}", "caption": "Frame at {clock} into the meeting"},
}
INDENT = Cm(0.63)


class ReportError(texts.Failure):
    """A report that could not be built; nothing was written."""


@dataclasses.dataclass
class ReportResult:
    output: Path
    language: str
    images: int
    sections: int
    size_bytes: int
    seconds: float
    template: Path | None
    cover: bool
    fields: list = dataclasses.field(default_factory=list)  # the template's fields that were filled
    contents: int = 0    # entries written in the template's tables of contents
    dropped: int = 0     # paragraphs with text of the template's model, left out


# ── The company's template ───────────────────────────────────────────────────

def _word_attribute(element, local):
    """The value of a WordprocessingML attribute of element, with or without a
    prefix, or None."""
    for name, value in element.attrib.items():
        namespace, _, found = name[1:].partition("}") if name.startswith("{") else ("", "", name)
        if found == local and (not namespace or namespace in WORD_NAMESPACES):
            return value
    return None


Field = collections.namedtuple("Field", "name text unnamed nested")


def _field(prefix, text, nested):
    """A field from the text of its instruction written before any field inside
    it (prefix), all the text it writes itself (text), and whether another field
    sits in its instruction. Its name is the first word of the prefix, which has
    to be written out whole: a prefix that is empty, or whose word touches the
    field inside it (IN{QUOTE "CLUDETEXT"}), leaves the field unnamed, since
    what Word would call it is not in the file. A formula (=) is named by its
    sign, and may go on with anything."""
    lead = prefix.lstrip()
    if not lead:
        return Field("", text, nested, nested)
    if lead.startswith("="):
        return Field("=", text, False, nested)
    word = lead.split(None, 1)[0]
    return Field(word, text, nested and lead == word, nested)


def _closed(open_field):
    text = "".join(open_field["pieces"])
    return _field(open_field["prefix"] if open_field["nested"] else text, text, open_field["nested"])


def _parse(data):
    """(the root of an XML part, {element: the prefixes declared at it}): the
    prefixes are kept only for the elements that use Markup Compatibility,
    whose attributes name prefixes."""
    root, scopes, levels, pending = None, {}, [], []
    for event, item in ElementTree.iterparse(io.BytesIO(data), events=("start-ns", "start", "end")):
        if event == "start-ns":
            pending.append(item)
        elif event == "start":
            root = item if root is None else root
            levels.append(dict(pending))
            pending = []
            if item.tag.startswith(f"{{{MC_NAMESPACE}}}") or any(key.startswith(f"{{{MC_NAMESPACE}}}") for key in item.attrib):
                scopes[item] = {prefix: uri for level in levels for prefix, uri in level.items()}
        else:
            levels.pop()
    return root, scopes


def _events(root):
    """("start" | "end", element) for the element and all below it, in document
    order, without recursion: a part may be nested deeper than the interpreter
    allows."""
    stack = [(root, iter(root))]
    yield "start", root
    while stack:
        element, children = stack[-1]
        child = next(children, None)
        if child is None:
            stack.pop()
            yield "end", element
        else:
            stack.append((child, iter(child)))
            yield "start", child


def _namespace(element):
    return element.tag[1:].partition("}")[0] if isinstance(element.tag, str) and element.tag.startswith("{") else ""


def _named_namespaces(element, attribute, scopes):
    """(the namespaces that the Markup Compatibility attribute of an element
    names by prefix, the prefixes it names that nothing declares)."""
    value = element.attrib.get(f"{{{MC_NAMESPACE}}}{attribute}")
    if value is None:
        return frozenset(), []
    scope, found, lost = scopes.get(element, {}), set(), []
    for item in value.split():
        prefix = item.split(":", 1)[0] if attribute in MC_NAME_ATTRIBUTES else item
        if prefix in scope:
            found.add(scope[prefix])
        else:
            lost.append(prefix)
    return frozenset(found), lost


def _is_branch(element):
    return _namespace(element) == MC_NAMESPACE and element.tag.rpartition("}")[2] in ("Choice", "Fallback")


def compatibility_problems(root, scopes):
    """What of Markup Compatibility in a part makes it impossible to know what
    Word would read, as (the message's key, the name of what it is about):
    an element or attribute of the namespace that is not known, an
    mc:AlternateContent that is not a list of mc:Choice, each with the
    namespaces it requires, and at most one mc:Fallback, last; a prefix that
    mc:Ignorable, mc:MustUnderstand, mc:ProcessContent or mc:Requires name and
    nothing declares; and a namespace of Word's own made ignorable."""
    problems, branches = set(), set()
    for element in root.iter():
        if _namespace(element) == MC_NAMESPACE:
            local = element.tag.rpartition("}")[2]
            children = [child.tag.rpartition("}")[2] if _namespace(child) == MC_NAMESPACE else "" for child in element]
            if local not in MC_ELEMENTS:
                problems.add(("report.active.compat_unknown", f"mc:{local}"))
            elif local == "AlternateContent":
                branches.update(id(child) for child in element)
                if "Choice" not in children or children.count("Fallback") > 1 or any(
                        kind not in ("Choice", "Fallback") for kind in children) or (
                        "Fallback" in children and children[-1] != "Fallback"):
                    problems.add(("report.active.compat_structure", "mc:AlternateContent"))
            elif id(element) not in branches:  # a Choice or a Fallback outside an mc:AlternateContent
                problems.add(("report.active.compat_structure", f"mc:{local}"))
            if local == "Choice":
                required = element.get("Requires", "").split()
                scope = scopes.get(element, {})
                if not required:
                    problems.add(("report.active.compat_structure", "mc:Choice"))
                problems.update(("report.active.compat_prefix", prefix) for prefix in required if prefix not in scope)
        for key in element.attrib:
            if key.startswith(f"{{{MC_NAMESPACE}}}"):
                local = key.rpartition("}")[2]
                if local not in MC_PREFIX_ATTRIBUTES | MC_NAME_ATTRIBUTES:
                    problems.add(("report.active.compat_unknown", f"mc:{local}"))
                    continue
                found, lost = _named_namespaces(element, local, scopes)
                problems.update(("report.active.compat_prefix", prefix) for prefix in lost)
                if local == "Ignorable" and found & (WORD_NAMESPACES | {MC_NAMESPACE}):
                    problems.add(("report.active.compat_prefix", "mc:Ignorable"))
    return sorted(problems)


def _finish(state, fields):
    """Close what a reading context left open, and read its loose text as one more field."""
    fields.extend(_closed(open_field) for open_field in state["open"])
    text = "".join(state["loose"])
    fields.append(_field(text, text, False))


def field_instructions(root, scopes=None):
    """Every field of an XML part, as Field, in document order: a simple field
    (w:fldSimple) by its w:instr, and a complex one by the w:instrText pieces
    joined from its begin up to its separate or end, across runs and
    paragraphs, each field on its own (a field inside another's instruction is
    its own, and its text is not the outer's). Text that no field holds, or that
    a field writes after its separate, is joined and read as one more field.

    Word reads one branch of an mc:AlternateContent, and skips an element of an
    ignorable namespace with what is inside it, so the pieces of two branches,
    or of an element Word skips and one it reads, are not one instruction: each
    mc:Choice, each mc:Fallback and each element in a namespace that mc:Ignorable
    names is a reading context of its own, whose fields are all judged (which
    one Word reads is not known). `scopes` is what _parse gives."""
    scopes = scopes or {}
    fields, states, isolated, ignorable = [], [{"open": [], "loose": []}], [], [frozenset()]
    for event, element in _events(root):
        if event == "end":
            ignorable.pop()
            if isolated and isolated[-1] is element:
                isolated.pop()
                _finish(states.pop(), fields)
            continue
        names = ignorable[-1] | _named_namespaces(element, "Ignorable", scopes)[0] if element in scopes else ignorable[-1]
        ignorable.append(names)
        if _is_branch(element) or _namespace(element) in names:
            isolated.append(element)
            states.append({"open": [], "loose": []})
        open_fields, loose = states[-1]["open"], states[-1]["loose"]
        kind = FIELD_TAGS.get(element.tag)
        if kind == "fldSimple":
            instruction = _word_attribute(element, "instr")
            if instruction is not None:
                fields.append(_field(instruction, instruction, False))
        elif kind in ("instrText", "delInstrText"):  # a tracked deletion comes back when it is rejected
            top = open_fields[-1] if open_fields else None
            (loose if top is None or top["separated"] else top["pieces"]).append("".join(element.itertext()))
        elif kind == "fldChar":
            step = _word_attribute(element, "fldCharType")
            if step == "begin":
                top = open_fields[-1] if open_fields else None
                if top is not None and not top["separated"] and not top["nested"]:
                    top["nested"], top["prefix"] = True, "".join(top["pieces"])
                open_fields.append({"pieces": [], "prefix": None, "nested": False, "separated": False})
            elif step == "separate" and open_fields:
                open_fields[-1]["separated"] = True
            elif step == "end" and open_fields:
                fields.append(_closed(open_fields.pop()))
    _finish(states[0], fields)
    return fields


def hyperlink_destination(field):
    """(where a HYPERLINK field goes, whether it names a place in the document
    with \\l): the first word or quoted text that is not a switch or a switch's
    argument; None when it has none."""
    rest = re.sub(r"^\s*\w+", "", field.text, count=1)
    destination, anchor, skip = None, False, False
    for match in re.finditer(r'"([^"]*)"|(\S+)', rest):
        value = match.group(1) if match.group(1) is not None else match.group(2)
        if skip:
            skip = False
        elif match.group(2) is not None and re.fullmatch(r"\\[A-Za-z]", value):
            skip = value.lower() in SWITCHES_WITH_ARGUMENT
            anchor = anchor or value.lower() == "\\l"
        elif destination is None:
            destination = value
    return destination, anchor


def _content_types(parts):
    """[(part name or extension, content type)] from [Content_Types].xml read
    as XML; empty when it cannot be read (active_content says so)."""
    try:
        root = ElementTree.fromstring(parts["[Content_Types].xml"])
    except (KeyError, ElementTree.ParseError, ValueError):
        return []
    return [(element.get("PartName") or "." + (element.get("Extension") or ""), element.get("ContentType") or "")
            for element in root.iter() if element.get("ContentType") is not None]


def carries_macros(parts):
    """Whether the package declares a macro-enabled content type, or holds a
    macro project (a part named vbaProject...). The content types are read as
    XML: however they are written, a content type is what the parser decodes."""
    return (any(Path(name).name.lower().startswith("vbaproject") for name in parts)
            or any(macro in content_type.lower() for _, content_type in _content_types(parts)
                   for macro in MACRO_TYPES))


def _is_xml(name, types):
    """Whether a part is XML: by its extension, or by a content type that says
    so (Word reads a part by its content type, not by its name)."""
    name = name.lower()
    if name.endswith(XML_EXTENSIONS):
        return True
    declared = {part.lower(): content_type for part, content_type in types}
    content_type = declared.get("/" + name) or declared.get(name) or declared.get(Path(name).suffix, "")
    return content_type.lower().endswith("xml")


def active_content(parts):
    """What in a Word package would be loaded or run from outside it when the
    document opens, one message each: an external relationship other than a
    hyperlink (an attached template, which may be a .dotm with macros; a
    linked picture), a relationship to an embedded object, control or macro
    project, a field that is not one of ALLOWED_FIELDS (only the field's name is
    compared, as written out in the file) or whose name another field builds,
    and a hyperlink to anything but a web page, a mail address or a place in
    the document. Every XML part is read as XML, with its namespaces, so the same
    instruction written in any equivalent way (a character reference, single
    quotes, another prefix, CDATA, a comment in the middle) is read alike;
    one that cannot be read is said, naming the part, since what it holds is
    not known."""
    found = []
    types = _content_types(parts)
    for name, data in sorted(parts.items()):
        if not _is_xml(name, types):
            continue
        try:
            root, scopes = _parse(data)
        except (ElementTree.ParseError, ValueError) as error:
            found.append(texts.Message("report.active.unreadable", part=name, detail=texts.External(str(error))))
            continue
        if name.lower().endswith(".rels"):
            for relationship in root:
                kind = relationship.get("Type", "").rsplit("/", 1)[-1]
                target = relationship.get("Target", "")
                if relationship.get("TargetMode", "").lower() == "external" and kind.lower() != "hyperlink":
                    found.append(texts.Message("report.active.external", part=name, kind=kind, target=target))
                elif kind.lower() in _ACTIVE_RELATIONSHIPS:
                    found.append(texts.Message("report.active.embedded", part=name, kind=kind, target=target))
                elif kind.lower() == "hyperlink" and not HYPERLINK_ALLOWED.match(target):
                    found.append(texts.Message("report.active.hyperlink", part=name, target=target))
        names, links, unnamed = set(), set(), False
        for field in field_instructions(root, scopes):
            unnamed = unnamed or field.unnamed
            word = re.match(r"\w+", field.name)
            # A name that does not start with a word (a quote, a dash) is judged as written, not as no name.
            word = word.group() if word else field.name.strip()
            # Upper-cased only when ASCII: "ſEQ" is not SEQ, though Python's upper() would make it so.
            word = word.upper() if word.isascii() else word
            if not word:  # nothing written, or the text no field holds: there is no name to judge
                continue
            if word not in ALLOWED_FIELDS:
                names.add(word[:NAME_SHOWN])
            elif word == "HYPERLINK" and not field.unnamed:
                destination, anchor = hyperlink_destination(field)
                if destination is None and field.nested and not anchor:
                    links.add("an address that another field builds")
                elif destination is not None and not HYPERLINK_ALLOWED.match(destination):
                    links.add(destination)
        for word in sorted(names):
            found.append(texts.Message("report.active.field", part=name, field=word))
        if unnamed:
            found.append(texts.Message("report.active.unnamed_field", part=name))
        for target in sorted(links):
            found.append(texts.Message("report.active.hyperlink", part=name, target=target))
        for key, what in compatibility_problems(root, scopes):
            found.append(texts.Message(key, part=name, name=what[:NAME_SHOWN]))
    return found


def template_bytes(path):
    """The template at path as a Word document's bytes, or ReportError saying
    why it cannot be used. A .dotx is given the document content type, which
    python-docx needs; nothing else in it changes."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in MACRO_EXTENSIONS:
        raise ReportError("report.macro_extension", name=path.name, suffix=suffix)
    if suffix not in TEMPLATE_EXTENSIONS:
        raise ReportError("report.not_word", name=path.name)
    try:
        parts = word_package.read_parts(path)
        names = list(parts)
        types = parts["[Content_Types].xml"]
    except word_package.PackageError as error:
        raise ReportError("report.template_too_big", name=path.name, reason=error.message) from None
    except (OSError, KeyError, zipfile.BadZipFile) as error:
        raise ReportError("report.cannot_open", name=path.name, detail=texts.External(str(error))) from None
    if carries_macros(parts):
        raise ReportError("report.macros", name=path.name)
    active = active_content(parts)
    if active:
        raise ReportError("report.active", name=path.name, items=texts.Joined(active, "\n  "),
                          allowed=", ".join(sorted(ALLOWED_FIELDS)))
    if TEMPLATE_TYPE in types:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            for name in names:
                data = parts[name]
                archive.writestr(name, data.replace(TEMPLATE_TYPE, DOCUMENT_TYPE) if name == "[Content_Types].xml"
                                 else data)
        data = buffer.getvalue()
    else:
        data = path.read_bytes()
    try:
        opened = docx.Document(io.BytesIO(data))
    except Exception as error:  # python-docx raises several kinds for a broken package
        raise ReportError("report.not_a_document", name=path.name, detail=texts.External(str(error))) from None
    try:
        layout.read_layout(opened)
    except layout.LayoutError as error:
        raise ReportError("report.template_unusable", name=path.name, error=error.message) from None
    return data


@dataclasses.dataclass
class TemplateInfo:
    """What the installation's template is, and what was understood of it."""
    path: Path
    name: str | None      # the file's name when it was set; None if it was set before that was recorded
    set_utc: str | None
    fields: list
    tables_of_contents: int
    start: str | None     # "index", "marker" or None
    dropped: int          # paragraphs with text of the template's model, dropped from every report


def stored_template(data_dir=None):
    """The installation's template, or None when there is none."""
    path = Path(data_dir or store.default_data_dir()) / TEMPLATE_NAME
    return path if path.is_file() else None


def template_info(data_dir=None):
    """The installation's template as TemplateInfo, or None when there is
    none; ReportError if it can no longer be used."""
    path = stored_template(data_dir)
    if path is None:
        return None
    found = layout.read_layout(docx.Document(io.BytesIO(template_bytes(path))))
    try:
        record = json.loads((path.parent / TEMPLATE_RECORD).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        record = {}
    name = record.get("name") if isinstance(record, dict) and isinstance(record.get("name"), str) else None
    set_utc = record.get("set_utc") if name and isinstance(record.get("set_utc"), str) else None
    return TemplateInfo(path, name, set_utc, found.fields, len(found.tocs), found.start_kind, found.dropped)


def set_template(path, data_dir=None, name=None):
    """Check the template and keep a copy of it in the data folder, with the
    name it was given (by default, its own file name)."""
    data = template_bytes(path)
    folder = store.check_data_dir(data_dir or store.default_data_dir())
    target = folder / TEMPLATE_NAME
    record = {"name": Path(name or Path(path).name).name[:200], "set_utc": store._now_utc()}
    with store.data_lock(folder):  # each file whole or not at all (WI20, R01)
        disk.write_bytes(target, data)
        disk.write_text(folder / TEMPLATE_RECORD, json.dumps(record, ensure_ascii=False) + "\n")
    return target


def remove_template(data_dir=None):
    """Remove the installation's template and its record; True if there was
    one."""
    folder = store.check_data_dir(data_dir or store.default_data_dir())
    with store.data_lock(folder):
        path = stored_template(folder)
        if path is None:
            return False
        path.unlink()
        (path.parent / TEMPLATE_RECORD).unlink(missing_ok=True)
    return True


def example_template():
    """The bytes of a template to start from (see layout.example_template)."""
    return layout.example_template()


# ── Reading the summary ──────────────────────────────────────────────────────

def plain(text):
    """Text without Markdown emphasis markers."""
    return re.sub(r"\*\*|__|`", "", text).replace("*", "").strip()


def summary_headings(text):
    return [plain(match.group(2)) for match in map(HEADING.match, text.splitlines()) if match]


def summary_language(text):
    """'es' or 'en': the language whose section names the summary uses most,
    or else the language of its words."""
    headings = set(summary_headings(text))
    counts = {language: len(headings & set(names + qa.HEADINGS[language]))
              for language, names in writer.SECTIONS.items()}
    best = max(counts, key=counts.get)
    return best if counts[best] else writer.detect_language(text)


def clock(hours, minutes, seconds):
    hours = int(hours)
    return f"{hours}:{minutes}:{seconds}" if hours else f"{int(minutes)}:{seconds}"


def cited_frames(text, frames_dir):
    """The frames the summary names, in order of first mention, as {name:
    path}. A missing frame, text that mentions a frame in any other way, or
    two frames named as a range (whose two ends nobody chose, INGOL D-181),
    is a ReportError naming every case."""
    frames, problems = {}, []
    for number, line in enumerate(text.splitlines(), start=1):
        found = writer.FRAME_RANGE.search(line)
        if found:
            problems.append(texts.Message("report.problem.range", line=number, text=found.group(0)[:80]))
        for match in FRAME_REF.finditer(line):
            name = match.group(1)
            if name not in frames:
                frames[name] = Path(frames_dir) / name
                if not frames[name].is_file():
                    problems.append(texts.Message("report.problem.missing", line=number, name=name,
                                                  folder=str(Path(frames_dir).resolve())))
        leftover = FRAME_REF.sub("", line)
        if FRAME_LIKE.search(leftover):
            problems.append(texts.Message("report.problem.unnamed", line=number, text=leftover.strip()[:80]))
    if problems:
        raise ReportError("report.not_built", problems=texts.Joined(problems, "\n  "))
    return frames


# ── Writing the document ─────────────────────────────────────────────────────

def _has_style(document, name):
    try:
        document.styles[name]
    except KeyError:
        return False
    return True


def _add_text(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1]).font.name = "Courier New"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _heading(document, text, level):
    """Add a heading; return its paragraph element."""
    if _has_style(document, f"Heading {level}"):
        return document.add_heading(text, level=level)._p
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt({1: 16, 2: 13}.get(level, 12))
    return paragraph._p


def _list_item(document, marker, text, depth):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = INDENT * (depth + 1)
    paragraph.paragraph_format.first_line_indent = -INDENT
    paragraph.add_run(f"{marker}\t")
    _add_text(paragraph, text)
    return paragraph


def _borders(table):
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "A0A0A0")
        borders.append(element)
    properties.append(borders)


def _table(document, lines, mention):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")]
            for line in lines if not TABLE_SEPARATOR.match(line)]
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    _borders(table)
    for r, row in enumerate(rows):
        for c in range(width):
            paragraph = table.rows[r].cells[c].paragraphs[0]
            _add_text(paragraph, mention(row[c]) if c < len(row) else "")
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if r == 0:
                    run.bold = True


def _picture(document, path, caption, width):
    holder = document.add_paragraph()
    holder.paragraph_format.keep_with_next = True
    holder.add_run().add_picture(str(path), width=width)
    run = document.add_paragraph().add_run(caption)
    run.italic = True
    run.font.size = Pt(9)


def _blocks(text):
    """The summary as (kind, payload) blocks: heading, bullet, numbered,
    table or paragraph. Blank lines and horizontal rules are dropped."""
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip() or RULE.match(line):
            i += 1
        elif (match := HEADING.match(line)):
            yield "heading", (max(1, min(len(match.group(1)) - 1, 3)), plain(match.group(2)))
            i += 1
        elif line.lstrip().startswith("|"):
            start = i
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                i += 1
            yield "table", lines[start:i]
        elif (match := BULLET.match(line)):
            yield "bullet", (len(match.group(1).expandtabs(4)) // 2, match.group(2))
            i += 1
        elif (match := NUMBERED.match(line)):
            yield "numbered", (len(match.group(1).expandtabs(4)) // 2, f"{match.group(2)}.", match.group(3))
            i += 1
        else:
            yield "paragraph", line.strip()
            i += 1


def node_text(element):
    """The text of a document element, from its w:t nodes."""
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def _body_is_empty(document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag != qn("w:p") or node_text(child).strip() or child.findall(".//" + qn("w:drawing")):
            return False
    return True


def _clear_body(document):
    body = document.element.body
    for child in list(body.iterchildren()):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _body_images(document):
    """The sha256 of every image drawn in the document's body, in order."""
    digests = []
    for blip in document.element.body.iter(qn("a:blip")):
        target = blip.get(qn("r:embed"))
        if target:
            digests.append(hashlib.sha256(document.part.related_parts[target].blob).hexdigest())
    return digests


def check_active_content(path):
    """ReportError unless the package at path is free of what a template is
    refused for (macros, content loaded or run from outside, a field that is
    not an allowed one, a part that is not readable XML). A template is checked when it
    is set, but a report is what the client opens, so it is checked again,
    read as it is written now, and not delivered if it fails."""
    try:
        parts = word_package.read_parts(path)
    except word_package.PackageError as error:
        raise ReportError("report.too_big", reason=error.message) from None
    found = [texts.Message("report.active.macros")] if carries_macros(parts) else []
    found += active_content(parts)
    if found:
        raise ReportError("report.active_in_report", items=texts.Joined(found, "\n  "),
                          allowed=", ".join(sorted(ALLOWED_FIELDS)))


def check_report(path, headings, cover_images, frames, *, start=0, contents=()):
    """ReportError unless the document at path has every summary heading, in
    order, after its first `start` body elements (the cover, whose table of
    contents repeats the headings); exactly the cover's images plus one of
    each cited frame; each table of contents of the cover listing exactly
    what `contents` says, in order; no known field left unfilled; and no
    active content (check_active_content)."""
    check_active_content(path)
    document = docx.Document(str(path))
    children = layout.body_children(document)
    texts = iter(layout.paragraph_text(child).strip() for child in children[start:] if child.tag == qn("w:p"))
    missing = [heading for heading in headings if not any(text == heading for text in texts)]
    if missing:
        raise ReportError("report.missing_sections", missing=missing)
    cover_ids = {id(child) for child in children[:start]}
    body = document.element.body
    found_contents = [layout.toc_entries(document, toc) for toc in layout._tocs(document)
                      if id(layout._top(toc[0], body)) in cover_ids]
    if found_contents != [list(entries) for entries in contents]:
        raise ReportError("report.toc_mismatch")
    left = layout.leftover_fields(document, children[:start])
    if left:
        raise ReportError("report.fields_left", fields=", ".join(left))
    expected = collections.Counter(cover_images)
    expected.update(hashlib.sha256(Path(frame).read_bytes()).hexdigest() for frame in frames)
    found = collections.Counter(_body_images(document))
    if found != expected:
        lacking = sum((expected - found).values())
        extra = sum((found - expected).values())
        raise ReportError("report.images_mismatch", lacking=lacking, extra=extra)


def build_report(frames_dir, *, title=None, date=None, project_name=None, data_dir=None, neutral=False,
                 template=None, client=None, meeting_type=None):
    """Write OUTPUT_NAME next to the summary in frames_dir and return what was
    built. `template` overrides the installation's; `neutral` uses none. The
    title, date, project, client and type fill the template's fields."""
    started = time.monotonic()
    frames_dir = Path(frames_dir)
    work_tree = store.enclosing_git_work_tree(frames_dir)
    if work_tree is not None:
        raise ReportError("report.inside_repository", folder=str(frames_dir.resolve()), work_tree=str(work_tree))
    summary = frames_dir / SUMMARY_NAME
    if not summary.is_file():
        raise ReportError("report.no_summary", summary=SUMMARY_NAME, folder=str(frames_dir.resolve()))
    # utf-8-sig: an editor may add a byte-order mark when the summary is
    # edited by hand, and it would hide the first heading.
    text = summary.read_text(encoding="utf-8-sig")
    headings = summary_headings(text)
    if not headings:
        raise ReportError("report.no_heading", path=str(summary.resolve()))
    unread = [number for number, line in enumerate(text.splitlines(), start=1)
              if line.lstrip().startswith("#") and not HEADING.match(line)]
    if unread:
        raise ReportError("report.unread_headings", lines=unread, path=str(summary.resolve()))
    frames = cited_frames(text, frames_dir)
    language = summary_language(text)
    labels = LABELS[language]

    template = None if neutral else (Path(template) if template else stored_template(data_dir))
    found = None
    if template is not None:
        document = docx.Document(io.BytesIO(template_bytes(template)))
        found = layout.read_layout(document)
        layout.drop_model(document, found)
        layout.fill_fields(document, layout.field_values(
            language, client=client, project=project_name, meeting=title.strip() if title else None, date=date,
            meeting_type=meeting_type))
        layout.no_update_on_open(document)
        cover = bool(found.tocs) or not _body_is_empty(document)
        if not cover:
            _clear_body(document)
    else:
        document = docx.Document()
        document.styles["Normal"].font.name = "Arial"
        document.styles["Normal"].font.size = Pt(11)
        for section in document.sections:
            section.top_margin = section.bottom_margin = Cm(2.5)
            section.left_margin = section.right_margin = Cm(2.5)
        cover = False
    cover_images = _body_images(document)
    if cover and not layout.ends_on_a_new_page(document):
        # A cover ending with its own section break already starts the report on a new page.
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    section = document.sections[-1]
    try:
        width = section.page_width - section.left_margin - section.right_margin
    except TypeError:
        width = Cm(16)

    heading_title = title.strip() if title and title.strip() else labels["title"]
    if _has_style(document, "Title"):
        first = document.add_paragraph(heading_title, style="Title")
    else:
        first = document.add_paragraph()
        run = first.add_run(heading_title)
        run.bold = True
        run.font.size = Pt(20)
    facts = [f"{labels['date']}: {date}" if date else "", f"{labels['project']}: {project_name}" if project_name else ""]
    if any(facts):
        document.add_paragraph(" · ".join(fact for fact in facts if fact))

    embedded, placed = set(), []

    def mention(line):
        found = [match.group(1) for match in FRAME_REF.finditer(line)]
        pending.extend(name for name in found if name not in embedded and name not in pending)
        return FRAME_REF.sub(lambda m: labels["mention"].format(clock=clock(*m.groups()[1:])), line)

    for kind, payload in _blocks(text):
        pending = []
        if kind == "heading":
            shown = mention(payload[1])
            placed.append((payload[0], shown, _heading(document, shown, payload[0])))
        elif kind == "table":
            _table(document, payload, mention)
        elif kind == "bullet":
            _list_item(document, "•", mention(payload[1]), payload[0])
        elif kind == "numbered":
            _list_item(document, payload[1], mention(payload[2]), payload[0])
        else:
            _add_text(document.add_paragraph(), mention(payload))
        for name in pending:
            match = FRAME_REF.match(f"[{name}]")
            try:
                _picture(document, frames[name], labels["caption"].format(clock=clock(*match.groups()[1:])), width)
            except (UnrecognizedImageError, OSError) as error:
                raise ReportError("report.frame_unembeddable", name=name,
                                  detail=texts.External(f"{type(error).__name__}: {error}")) from None
            embedded.add(name)

    expected = [FRAME_REF.sub(lambda m: labels["mention"].format(clock=clock(*m.groups()[1:])), heading)
                for heading in headings]
    contents = []
    if found is not None and found.tocs:
        bookmarks = layout.bookmark_headings(document, placed)
        try:
            for toc in found.tocs:
                layout.fill_toc(document, toc, placed, bookmarks)
        except layout.LayoutError as error:
            raise ReportError("report.template_unusable", name=Path(template).name, error=error.message) from None
        # What each table of contents must list, from the summary's headings
        # and the field's levels, not from what was written in it.
        depth = [max(1, min(len(HEADING.match(line).group(1)) - 1, 3)) for line in text.splitlines()
                 if HEADING.match(line)]
        for _, _, instruction in found.tocs:
            low, high = layout.levels(instruction)
            contents.append([heading for level, heading in zip(depth, expected) if low <= level <= high])
    # Where the report starts, counted after the tables of contents were
    # rewritten: they may have changed how many paragraphs the cover has.
    start = layout.body_children(document).index(first._p)

    output = frames_dir / OUTPUT_NAME
    partial = frames_dir / (OUTPUT_NAME + ".partial")
    try:
        document.save(str(partial))
        check_report(partial, expected, cover_images, frames.values(), start=start, contents=contents)
    except OSError as error:
        partial.unlink(missing_ok=True)
        raise ReportError("report.cannot_write", folder=str(frames_dir.resolve()),
                          detail=texts.External(str(error))) from None
    except BaseException:
        partial.unlink(missing_ok=True)
        raise
    try:
        os.replace(partial, output)
    except OSError as error:
        partial.unlink(missing_ok=True)
        hint = texts.Message("report.close_word") if isinstance(error, PermissionError) else ""
        raise ReportError("report.cannot_replace", path=str(output.resolve()),
                          detail=texts.External(str(error.strerror)), hint=hint) from None
    return ReportResult(output, language, len(frames), len(headings), output.stat().st_size,
                        time.monotonic() - started, template, cover, found.fields if found else [],
                        sum(len(entries) for entries in contents), found.dropped if found else 0)
