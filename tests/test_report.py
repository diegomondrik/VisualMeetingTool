"""Tests for meetingtool.report. Summaries, frames and company templates are
invented at test time in a temporary folder: no Word file is committed, and
nothing here reaches the network or reads a key."""

import contextlib
import hashlib
import io
import socket
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

import docx
import docx.text.paragraph
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image

from meetingtool.projects import store
from meetingtool.report import document, layout
from meetingtool.report.__main__ import main
from meetingtool.summary import writer

REPOSITORY = Path(__file__).resolve().parent.parent
FRAMES = {"frame_001_t00-01-22.jpg": (200, 30, 30), "frame_002_t00-13-03.jpg": (30, 30, 200),
          "frame_003_t01-02-03.jpg": (30, 160, 30)}
SCREEN = writer.SECTIONS["es"][4]


def summary_text(language="es", screen=None):
    sections = writer.required_headings(language)
    bodies = {name: f"Texto de **{name}** con *énfasis*." for name in sections}
    bodies[sections[1]] = "| Nombre | Empresa | Rol |\n|---|---|---|\n| Ana | ACME | **Líder** |"
    bodies[sections[4]] = screen if screen is not None else (
        "* **[frame_001_t00-01-22.jpg]** y **[frame_002_t00-13-03.jpg]**: la planilla de costos.\n"
        "* **[frame_002_t00-13-03.jpg]**: la misma planilla, otra vez.")
    bodies[sections[3]] = "1. Mandar el detalle\n2. Revisar el total\n  - con el equipo"
    return "\n\n".join(f"## {name}\n{bodies[name]}" for name in sections) + "\n"


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def company_template(path, *, cover="Portada de Empresa Inventada", logo=None, word_like=False):
    """A company's template: header with a logo and name, footer, coloured
    headings, its own body font, and a cover page. word_like removes the list
    and table styles, as in a template saved by Word."""
    template = docx.Document()
    header = template.sections[0].header.paragraphs[0]
    if logo:
        header.add_run().add_picture(str(logo))
    header.add_run("Empresa Inventada SA")
    template.sections[0].footer.paragraphs[0].text = "www.inventada.example"
    template.styles["Heading 1"].font.color.rgb = RGBColor(0xC0, 0x10, 0x10)
    template.styles["Normal"].font.name = "Georgia"
    template.styles["Normal"].font.size = Pt(10)
    if cover:
        template.add_paragraph(cover)
    if word_like:
        styles = template.styles.element
        for style in list(styles.iterchildren(qn("w:style"))):
            if style.get(qn("w:styleId")) in {"ListBullet", "ListNumber", "TableGrid", "Caption"}:
                styles.remove(style)
    template.save(str(path))
    if path.suffix.lower() in (".dotx", ".dotm"):
        as_template(path, path)
    return path


def as_template(source, target, content_type=document.TEMPLATE_TYPE):
    """Rewrite a .docx package with another main content type (a .dotx, or a
    macro-enabled one)."""
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(document.DOCUMENT_TYPE, content_type)
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


class Workspace(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.frames = self.tmp / "frames-out"
        self.frames.mkdir()
        for name, colour in FRAMES.items():
            Image.new("RGB", (320, 180), colour).save(self.frames / name)
        self.data = self.tmp / "data"
        self.write_summary(summary_text())

    def tearDown(self):
        self._tmp.cleanup()

    def write_summary(self, text):
        (self.frames / document.SUMMARY_NAME).write_text(text, encoding="utf-8")

    def build(self, **kwargs):
        kwargs.setdefault("data_dir", self.data)
        return document.build_report(self.frames, **kwargs)

    def output(self):
        return self.frames / document.OUTPUT_NAME

    def assertNothingWritten(self):
        self.assertFalse(self.output().exists())
        self.assertEqual(list(self.frames.glob("*.partial")), [])

    @staticmethod
    def texts(path):
        return [paragraph.text for paragraph in docx.Document(str(path)).paragraphs]


class BuildTest(Workspace):
    """WI11-AC01."""

    def test_the_summary_becomes_a_word_document_next_to_it(self):
        result = self.build(title="Reunión de prueba", date="2026-09-10", project_name="Planta")
        self.assertEqual(result.output, self.output())
        self.assertEqual(result.sections, len(writer.required_headings("es")))
        report = docx.Document(str(self.output()))
        headings = [p.text for p in report.paragraphs if p.style.name == "Heading 1"]
        self.assertEqual(headings, writer.required_headings("es"))
        texts = [p.text for p in report.paragraphs]
        self.assertEqual(texts[:2], ["Reunión de prueba", "Fecha: 2026-09-10 · Proyecto: Planta"])
        table = report.tables[0]
        self.assertEqual([c.text for c in table.rows[1].cells], ["Ana", "ACME", "Líder"])
        self.assertTrue(all(run.bold for run in table.rows[0].cells[0].paragraphs[0].runs))
        bold = [run.text for p in report.paragraphs for run in p.runs if run.bold]
        self.assertIn("Resumen ejecutivo", bold)
        italic = [run.text for p in report.paragraphs for run in p.runs if run.italic]
        self.assertIn("énfasis", italic)
        self.assertIn("1.\tMandar el detalle", texts)
        self.assertIn("•\tcon el equipo", texts)
        self.assertEqual(list(self.frames.glob("*.partial")), [])

    def test_it_can_be_built_again_after_the_summary_changes(self):
        self.build()
        self.write_summary(summary_text().replace("con *énfasis*", "corregido a mano"))
        self.build()
        self.assertTrue(any("corregido a mano" in text for text in self.texts(self.output())))

    def test_a_frames_folder_inside_a_git_work_tree_is_refused_before_writing(self):
        (self.tmp / ".git").mkdir()
        with self.assertRaisesRegex(document.ReportError, "inside the git work tree"):
            self.build()
        self.assertNothingWritten()

    def test_without_a_summary_nothing_is_built(self):
        (self.frames / document.SUMMARY_NAME).unlink()
        with self.assertRaisesRegex(document.ReportError, "write the summary first"):
            self.build()
        self.assertNothingWritten()


class ImagesTest(Workspace):
    """WI11-AC02."""

    def test_each_mentioned_frame_is_embedded_once_after_its_first_mention(self):
        result = self.build()
        self.assertEqual(result.images, 2)
        report = docx.Document(str(self.output()))
        self.assertEqual(sorted(document._body_images(report)),
                         sorted(digest(self.frames / name) for name in list(FRAMES)[:2]))
        body = [(child.tag, document.node_text(child)) for child in report.element.body.iterchildren()]
        drawings = [i for i, child in enumerate(report.element.body.iterchildren())
                    if child.findall(".//" + qn("w:drawing"))]
        first_mention = next(i for i, (_, text) in enumerate(body) if "imagen 1:22 y imagen 13:03" in text)
        self.assertEqual(drawings, [first_mention + 1, first_mention + 3])
        self.assertEqual(body[first_mention + 2][1], "Imagen del minuto 1:22 de la reunión")
        self.assertEqual(body[first_mention + 4][1], "Imagen del minuto 13:03 de la reunión")
        self.assertIn("imagen 13:03: la misma planilla, otra vez.", body[first_mention + 5][1])

    def test_a_frame_the_summary_does_not_mention_is_left_out(self):
        self.build()
        report = docx.Document(str(self.output()))
        self.assertNotIn(digest(self.frames / "frame_003_t01-02-03.jpg"), document._body_images(report))

    def test_a_frame_mentioned_in_a_table_follows_the_table(self):
        self.write_summary(summary_text(screen="| Imagen | Qué muestra |\n|---|---|\n"
                                               "| [frame_003_t01-02-03.jpg] | el tablero |"))
        self.build()
        report = docx.Document(str(self.output()))
        self.assertEqual(report.tables[1].rows[1].cells[0].text, "imagen 1:02:03")
        self.assertEqual(document._body_images(report), [digest(self.frames / "frame_003_t01-02-03.jpg")])

    def test_an_english_summary_gets_english_labels(self):
        self.write_summary(summary_text("en", screen="- [frame_002_t00-13-03.jpg]: the cost sheet"))
        result = self.build(date="2026-09-10", project_name="Plant")
        self.assertEqual(result.language, "en")
        texts = self.texts(self.output())
        self.assertEqual(texts[:2], ["Meeting summary", "Date: 2026-09-10 · Project: Plant"])
        self.assertIn("Frame at 13:03 into the meeting", texts)
        self.assertIn("•\tframe 13:03: the cost sheet", texts)


class RefusalTest(Workspace):
    """WI11-AC03."""

    def test_a_missing_frame_stops_the_build_and_is_named(self):
        (self.frames / "frame_002_t00-13-03.jpg").unlink()
        with self.assertRaisesRegex(document.ReportError, r"frame_002_t00-13-03\.jpg is not in"):
            self.build()
        self.assertNothingWritten()

    def test_a_mention_that_names_no_frame_file_stops_the_build(self):
        for mention in ("[frame_017, t00:13:03]", "[frames_019–024, t00:20:00-00:24:00]", "(Frame_001_t00-01-22)",
                        "[frame_001_t00-01-22.png]"):
            with self.subTest(mention=mention):
                self.write_summary(summary_text(screen=f"- {mention}: la planilla"))
                with self.assertRaisesRegex(document.ReportError, "names no frame file"):
                    self.build()
                self.assertNothingWritten()

    def test_a_range_of_frames_stops_the_build(self):
        # INGOL D-181: the report showed the two ends of the range, which
        # nobody chose (the real one began on rows without data).
        self.write_summary(summary_text(screen="- `[frame_001_t00-01-22.jpg]` a `[frame_002_t00-13-03.jpg]`: "
                                               "la planilla"))
        with self.assertRaisesRegex(document.ReportError, "a range of frames; name each frame on its own"):
            self.build()
        self.assertNothingWritten()

    def test_every_problem_is_named_at_once(self):
        self.write_summary(summary_text(screen="- [frame_009_t00-09-09.jpg] y [frame_017, t00:13:03]"))
        with self.assertRaises(document.ReportError) as raised:
            self.build()
        self.assertIn("frame_009_t00-09-09.jpg is not in", str(raised.exception))
        self.assertIn("names no frame file", str(raised.exception))

    def test_a_failed_build_leaves_the_earlier_report_as_it_was(self):
        self.build()
        before = digest(self.output())
        (self.frames / "frame_001_t00-01-22.jpg").unlink()
        with self.assertRaises(document.ReportError):
            self.build()
        self.assertEqual(digest(self.output()), before)


class CompletenessTest(Workspace):
    """WI11-AC04: the check fails when the document lacks what the summary has."""

    def test_a_document_missing_an_image_is_not_delivered(self):
        real, calls = document._picture, []

        def skip_second(*args):
            calls.append(args)
            if len(calls) != 2:
                real(*args)

        with mock.patch.object(document, "_picture", skip_second), \
                self.assertRaisesRegex(document.ReportError, "1 image\\(s\\) missing"):
            self.build()
        self.assertNothingWritten()

    def test_a_document_missing_a_section_is_not_delivered(self):
        real = document._heading

        def skip(doc, text, level):
            if text != SCREEN:
                real(doc, text, level)

        with mock.patch.object(document, "_heading", skip), \
                self.assertRaisesRegex(document.ReportError, "missing the section"):
            self.build()
        self.assertNothingWritten()

    def test_an_extra_image_is_not_delivered_either(self):
        real = document._picture

        def twice(doc, path, caption, width):
            real(doc, path, caption, width)
            real(doc, self.frames / "frame_003_t01-02-03.jpg", caption, width)

        with mock.patch.object(document, "_picture", twice), \
                self.assertRaisesRegex(document.ReportError, "unexpected"):
            self.build()
        self.assertNothingWritten()

    def test_sections_out_of_order_are_not_delivered(self):
        self.build()
        headings = writer.required_headings("es")
        swapped = [headings[1], headings[0]] + headings[2:]
        frames = [self.frames / name for name in list(FRAMES)[:2]]
        with self.assertRaisesRegex(document.ReportError, "missing the section"):
            document.check_report(self.output(), swapped, [], frames)
        document.check_report(self.output(), headings, [], frames)


class NoNetworkNoKeyTest(Workspace):
    """WI11-AC05."""

    def test_building_opens_no_network_connection(self):
        def refuse(*args, **kwargs):
            raise AssertionError("the report tried to open a network connection")

        template = company_template(self.tmp / "company.docx")
        with mock.patch.object(socket, "socket", refuse), mock.patch.object(socket, "create_connection", refuse):
            self.build()
            self.build(template=template)
        self.assertTrue(self.output().is_file())

    def test_the_report_command_never_loads_the_key_store(self):
        probe = ("import sys\nimport meetingtool.report.__main__\n"
                 "print('meetingtool.reading.credentials' in sys.modules)\n")
        result = subprocess.run([sys.executable, "-c", probe], cwd=REPOSITORY, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout.strip(), "False")


class TemplateTest(Workspace):
    """WI11-AC06."""

    def setUp(self):
        super().setUp()
        self.logo = self.tmp / "logo.png"
        Image.new("RGB", (120, 40), (250, 200, 0)).save(self.logo)

    def check_company_design(self, report_path, cover="Portada de Empresa Inventada"):
        report = docx.Document(str(report_path))
        section = report.sections[0]
        self.assertEqual(section.header.paragraphs[0].text, "Empresa Inventada SA")
        self.assertEqual(section.footer.paragraphs[0].text, "www.inventada.example")
        logos = [rel for rel in section.header.part.rels.values() if "image" in rel.reltype]
        self.assertEqual(len(logos), 1)
        self.assertEqual(logos[0].target_part.blob, self.logo.read_bytes())
        self.assertEqual(report.styles["Heading 1"].font.color.rgb, RGBColor(0xC0, 0x10, 0x10))
        self.assertEqual(report.styles["Normal"].font.name, "Georgia")
        self.assertEqual(report.paragraphs[0].text, cover)
        page_break = report.paragraphs[1].runs[0]._r.find(qn("w:br"))
        self.assertEqual(page_break.get(qn("w:type")), "page")
        return report

    def test_a_docx_and_a_dotx_template_give_the_report_the_company_design(self):
        for name in ("company.docx", "company.dotx"):
            with self.subTest(template=name):
                document.set_template(company_template(self.tmp / name, logo=self.logo), self.data)
                result = self.build()
                self.assertTrue(result.cover)
                self.assertEqual(result.template, self.data / document.TEMPLATE_NAME)
                report = self.check_company_design(self.output())
                self.assertEqual(sorted(document._body_images(report)),
                                 sorted(digest(self.frames / n) for n in list(FRAMES)[:2]))

    def test_a_template_saved_by_word_without_list_or_table_styles_works(self):
        document.set_template(company_template(self.tmp / "word.docx", logo=self.logo, word_like=True), self.data)
        template = docx.Document(str(self.data / document.TEMPLATE_NAME))
        for missing in ("List Bullet", "Table Grid"):
            with self.assertRaises(KeyError):
                template.styles[missing]
        self.build()
        report = self.check_company_design(self.output())
        self.assertEqual([c.text for c in report.tables[0].rows[1].cells], ["Ana", "ACME", "Líder"])

    def test_a_template_with_an_empty_page_gives_no_cover(self):
        document.set_template(company_template(self.tmp / "plain.docx", cover=None, logo=self.logo), self.data)
        result = self.build(title="Reunión de prueba")
        self.assertFalse(result.cover)
        texts = self.texts(self.output())
        self.assertEqual(texts[0], "Reunión de prueba")
        self.assertFalse(docx.Document(str(self.output())).element.body.findall(".//" + qn("w:br")))

    def test_the_cover_images_are_kept_and_counted_apart_from_the_frames(self):
        path = company_template(self.tmp / "cover.docx", logo=self.logo)
        with_logo = docx.Document(str(path))
        with_logo.paragraphs[0].add_run().add_picture(str(self.logo))
        with_logo.save(str(path))
        document.set_template(path, self.data)
        self.build()
        images = document._body_images(docx.Document(str(self.output())))
        self.assertEqual(images[0], digest(self.logo))
        self.assertEqual(len(images), 3)

    def test_without_a_template_or_with_neutral_the_design_is_neutral(self):
        self.build()
        neutral = docx.Document(str(self.output()))
        self.assertEqual(neutral.sections[0].header.paragraphs[0].text, "")
        self.assertEqual(neutral.styles["Normal"].font.name, "Arial")
        document.set_template(company_template(self.tmp / "company.docx", logo=self.logo), self.data)
        result = self.build(neutral=True)
        self.assertIsNone(result.template)
        self.assertEqual(docx.Document(str(self.output())).sections[0].header.paragraphs[0].text, "")

    def test_remove_goes_back_to_the_neutral_design(self):
        document.set_template(company_template(self.tmp / "company.docx"), self.data)
        self.assertTrue(document.remove_template(self.data))
        self.assertFalse(document.remove_template(self.data))
        self.assertIsNone(self.build().template)


class TemplateRefusalTest(Workspace):
    """WI11-AC07."""

    def assertRefused(self, path, pattern):
        with self.assertRaisesRegex(document.ReportError, pattern):
            document.set_template(path, self.data)

    def test_a_template_that_can_carry_macros_is_refused(self):
        source = company_template(self.tmp / "company.docx")
        for name in ("company.docm", "company.dotm"):
            with self.subTest(name=name):
                target = self.tmp / name
                target.write_bytes(source.read_bytes())
                self.assertRefused(target, "can carry macros")
        enabled = self.tmp / "renamed.docx"
        as_template(source, enabled, b"application/vnd.ms-word.document.macroEnabled.main+xml")
        self.assertRefused(enabled, "carries macros")
        with_project = self.tmp / "with-project.docx"
        with zipfile.ZipFile(source) as archive, zipfile.ZipFile(with_project, "w") as out:
            for item in archive.infolist():
                out.writestr(item, archive.read(item.filename))
            out.writestr("word/vbaProject.bin", b"\x00" * 16)
        self.assertRefused(with_project, "carries macros")
        self.assertIsNone(document.stored_template(self.data))

    def test_a_file_that_is_not_a_word_template_is_refused_when_set(self):
        text = self.tmp / "notes.docx"
        text.write_text("not a Word file", encoding="utf-8")
        self.assertRefused(text, "cannot be opened")
        broken = self.tmp / "broken.docx"
        with zipfile.ZipFile(broken, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
        self.assertRefused(broken, "cannot be opened")
        pdf = self.tmp / "design.pdf"
        pdf.write_bytes(b"%PDF-1.4")
        self.assertRefused(pdf, "not a Word document or template")
        self.assertRefused(self.tmp / "missing.docx", "cannot be opened")
        self.assertIsNone(document.stored_template(self.data))

    def test_a_refused_template_leaves_the_kept_one_in_place(self):
        kept = document.set_template(company_template(self.tmp / "company.docx"), self.data)
        before = digest(kept)
        (self.tmp / "bad.dotm").write_bytes(b"x")
        self.assertRefused(self.tmp / "bad.dotm", "can carry macros")
        self.assertEqual(digest(kept), before)

    def test_a_data_folder_inside_a_git_work_tree_is_refused(self):
        (self.tmp / ".git").mkdir()
        with self.assertRaises(store.ProjectError):
            document.set_template(company_template(self.tmp / "company.docx"), self.data)


class CommandLineTest(Workspace):
    def run_main(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = main(["--data-dir", str(self.data), *args])
        return code, out.getvalue(), err.getvalue()

    def test_template_set_show_remove_and_build(self):
        code, out, _ = self.run_main("template", "show")
        self.assertEqual((code, out.strip()), (0, "no template: reports use the neutral design"))
        code, out, _ = self.run_main("template", "set", str(company_template(self.tmp / "company.dotx")))
        self.assertEqual(code, 0)
        self.assertIn(document.TEMPLATE_NAME, out)
        code, out, _ = self.run_main("build", "--frames", str(self.frames), "--date", "2026-09-10")
        self.assertEqual(code, 0)
        self.assertIn("2 image(s)", out)
        self.assertIn("with its cover", out)
        code, out, _ = self.run_main("template", "remove")
        self.assertEqual((code, out.strip()), (0, "template removed"))

    def test_the_project_name_comes_from_the_project(self):
        project = store.create_project(self.data, "Planta Norte", "ACME")
        code, _, _ = self.run_main("build", "--frames", str(self.frames), "--project", project["id"])
        self.assertEqual(code, 0)
        self.assertIn("Proyecto: Planta Norte", self.texts(self.output()))
        code, _, err = self.run_main("build", "--frames", str(self.frames), "--project", "no-such-project")
        self.assertEqual(code, 2)
        self.assertIn("there is no project", err)

    def test_an_error_is_printed_and_nothing_is_written(self):
        (self.frames / "frame_001_t00-01-22.jpg").unlink()
        code, _, err = self.run_main("build", "--frames", str(self.frames))
        self.assertEqual(code, 2)
        self.assertIn("frame_001_t00-01-22.jpg is not in", err)
        self.assertNothingWritten()


RELATIONSHIPS = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">%s</Relationships>'
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"


def edit_package(source, target, *, add=None, insert=None):
    """Copy a Word package, adding parts (add: name -> bytes) and inserting
    bytes before a marker in a part (insert: name -> (marker, bytes))."""
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for name, (marker, data) in (insert or {}).items():
        parts[name] = parts[name].replace(marker, data + marker, 1)
    parts.update(add or {})
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return target


def relationship(kind, target, external=True):
    mode = ' TargetMode="External"' if external else ""
    return f'<Relationship Id="rId900" Type="{REL}{kind}" Target="{target}"{mode}/>'.encode()


def package_parts(path):
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


class ReviewCorrectionsTest(Workspace):
    """The independent review's findings on 3385859: P1-1 (content loaded
    from outside the template), P2-1, P2-2, P3-1, P3-3 and P3-4."""

    def setUp(self):
        super().setUp()
        self.clean = company_template(self.tmp / "clean.docx")

    def variant(self, name, **changes):
        return edit_package(self.clean, self.tmp / name, **changes)

    def test_p1_1_a_template_that_loads_or_runs_outside_content_is_refused(self):
        end = b"</Relationships>"
        variants = {
            "attached template": self.variant("attached.docx", add={
                "word/_rels/settings.xml.rels": RELATIONSHIPS % relationship(
                    "attachedTemplate", "file:///C:/elsewhere/Company.dotm")}),
            "embedded object": self.variant("ole.docx", add={"word/embeddings/oleObject1.bin": b"\x00" * 32},
                                            insert={"word/_rels/document.xml.rels": (end, relationship(
                                                "oleObject", "embeddings/oleObject1.bin", external=False))}),
            "linked picture": self.variant("linked.docx", insert={"word/_rels/document.xml.rels": (end, relationship(
                "image", "http://example.invalid/pixel.png"))}),
            "field split across runs": self.variant("field.docx", insert={"word/document.xml": (
                b"<w:sectPr", b'<w:p><w:r><w:instrText xml:space="preserve"> INCLUDEPIC</w:instrText></w:r>'
                b'<w:r><w:instrText>TURE "http://example.invalid/x.png"</w:instrText></w:r></w:p>')}),
        }
        for label, path in variants.items():
            with self.subTest(variant=label):
                with self.assertRaisesRegex(document.ReportError, "load or run from outside"):
                    document.set_template(path, self.data)
                self.assertIsNone(document.stored_template(self.data))

    def test_p1_1_dde_fields_are_found_even_when_split(self):
        # Checked in memory: a file with a DDE field may be locked by the antivirus.
        split = (b'<w:r><w:instrText xml:space="preserve"> DD</w:instrText></w:r>'
                 b'<w:r><w:instrText>EAUTO x y</w:instrText></w:r>')
        self.assertEqual(document.active_content({"word/document.xml": split}), ["word/document.xml: a DDEAUTO field"])
        simple = b'<w:fldSimple w:instr=" DDE x y"/>'
        self.assertEqual(document.active_content({"word/header1.xml": simple}), ["word/header1.xml: a DDE field"])
        harmless = b'<w:instrText> PAGE </w:instrText><w:fldSimple w:instr=" HYPERLINK &quot;x&quot;"/>'
        self.assertEqual(document.active_content({"word/footer1.xml": harmless}), [])

    def test_p1_1_a_hyperlink_and_a_template_saved_normally_are_accepted(self):
        linked = self.variant("hyperlink.docx", insert={"word/_rels/document.xml.rels": (
            b"</Relationships>", relationship("hyperlink", "https://www.inventada.example"))})
        document.set_template(linked, self.data)
        self.assertEqual(document.active_content(package_parts(self.clean)), [])

    def test_p1_1_a_template_placed_by_hand_is_checked_again_when_a_report_is_built(self):
        self.data.mkdir()
        self.variant("by-hand.docx", add={"word/_rels/settings.xml.rels": RELATIONSHIPS % relationship(
            "attachedTemplate", "file:///C:/elsewhere/Company.dotm")})
        (self.tmp / "by-hand.docx").replace(self.data / document.TEMPLATE_NAME)
        with self.assertRaisesRegex(document.ReportError, "load or run from outside"):
            self.build()
        self.assertNothingWritten()

    def test_p2_1_a_frame_named_in_a_heading_follows_the_heading(self):
        self.write_summary(summary_text(screen="### [frame_003_t01-02-03.jpg] Tablero de costos\nEl total."))
        result = self.build()
        self.assertEqual(result.images, 1)
        body = list(docx.Document(str(self.output())).element.body.iterchildren())
        heading = next(i for i, child in enumerate(body)
                       if document.node_text(child) == "imagen 1:02:03 Tablero de costos")
        self.assertTrue(body[heading + 1].findall(".//" + qn("w:drawing")))

    def test_p2_2_a_byte_order_mark_and_a_heading_without_a_space_are_still_sections(self):
        text = "\ufeff" + summary_text().replace("## Decisiones", "##Decisiones")
        (self.frames / document.SUMMARY_NAME).write_text(text, encoding="utf-8")
        result = self.build()
        self.assertEqual(result.sections, len(writer.required_headings("es")))
        headings = [p.text for p in docx.Document(str(self.output())).paragraphs if p.style.name == "Heading 1"]
        self.assertEqual(headings, writer.required_headings("es"))

    def test_p2_2_a_line_that_starts_with_a_hash_but_is_not_read_as_a_heading_stops_the_build(self):
        self.write_summary(summary_text().replace("## Decisiones", "   ## Decisiones"))
        with self.assertRaisesRegex(document.ReportError, "start with # but are not headings"):
            self.build()
        self.assertNothingWritten()

    def test_p3_1_a_word_that_ends_in_frame_is_not_a_frame_mention(self):
        self.write_summary(summary_text().replace("con *énfasis*", "con sales_dataframe_v2 y keyframe_interval"))
        self.assertEqual(self.build().images, 2)

    def test_p3_3_a_frame_that_is_not_an_image_is_named(self):
        (self.frames / "frame_002_t00-13-03.jpg").write_text("not an image", encoding="utf-8")
        with self.assertRaisesRegex(document.ReportError, "frame_002_t00-13-03.jpg cannot be embedded"):
            self.build()
        self.assertNothingWritten()

    def test_p3_4_a_report_open_in_word_is_named_and_the_earlier_one_stays(self):
        self.build()
        before = digest(self.output())
        with mock.patch.object(document.os, "replace", side_effect=PermissionError(13, "Permission denied")), \
                self.assertRaisesRegex(document.ReportError, "close it and try again"):
            self.build()
        self.assertEqual(digest(self.output()), before)
        self.assertEqual(list(self.frames.glob("*.partial")), [])


# ── WI16: the template's fields, where the report starts, its table of contents ──

MODEL = ["RESUMEN EJECUTIVO", "Texto Normal", "Caso 1: Erosión del Poder de Compra", "Organización - fecha"]
MEETING = {"title": "Sesión de dudas", "date": "2026-09-22", "project_name": "Planta Norte", "client": "ACME",
           "meeting_type": "requirements"}


def _run_of(node):
    run = OxmlElement("w:r")
    run.append(node)
    return run


def _field_char(kind, dirty=False):
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), kind)
    if dirty:
        node.set(qn("w:dirty"), "true")
    return node


def _instruction(text):
    node = OxmlElement("w:instrText")
    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    node.text = text
    return node


def _text(text):
    node = OxmlElement("w:t")
    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    node.text = text
    return node


def add_toc(template, stale, *, levels="1-3", sdt=False, dirty=False):
    """A table of contents as Word saves it: a TOC field over several
    paragraphs holding the entries Word last computed, the first with its own
    PAGEREF field; sdt wraps it in a content control, as Word's automatic
    table does."""
    made = []
    for number, entry in enumerate(stale):
        paragraph = template.add_paragraph(style="toc 1")._p
        if number == 0:
            for node in (_field_char("begin", dirty), _instruction(f' TOC \\o "{levels}" \\h \\z \\u '),
                         _field_char("separate")):
                paragraph.append(_run_of(node))
        paragraph.append(_run_of(_text(entry)))
        if number == 0:
            for node in (_field_char("begin"), _instruction(" PAGEREF _Toc100 \\h "), _field_char("separate"),
                         _text("3"), _field_char("end")):
                paragraph.append(_run_of(node))
        made.append(paragraph)
    made[-1].append(_run_of(_field_char("end")))
    if sdt:
        control = OxmlElement("w:sdt")
        properties = OxmlElement("w:sdtPr")
        part = OxmlElement("w:docPartObj")
        gallery = OxmlElement("w:docPartGallery")
        gallery.set(qn("w:val"), "Table of Contents")
        part.append(gallery)
        properties.append(part)
        content = OxmlElement("w:sdtContent")
        control.append(properties)
        control.append(content)
        made[0].addprevious(control)
        for paragraph in made:
            content.append(paragraph)
    return made


def split_field(paragraph, pieces, formats):
    """Write a paragraph as Word may keep it after an edit: one run per piece,
    each with its own format (bold, italic or none)."""
    for piece, style in zip(pieces, formats):
        run = paragraph.add_run(piece)
        run.bold = style == "bold"
        run.italic = style == "italic"


def owner_shaped(path, *, sdt=False, levels="1-3", dirty=False, update_fields=False):
    """A template shaped like the owner's (2026-09-30): a cover with fields,
    one split by Word, a table of contents holding what Word last saved, and
    after it a page of model with headings; fields in the header and footer."""
    template = docx.Document()
    for level in (1, 2, 3):
        template.styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH)
    template.sections[0].header.paragraphs[0].text = "{proyecto} · Empresa Inventada"
    template.sections[0].footer.paragraphs[0].text = "{fecha}"
    template.add_paragraph("{cliente}", style="Title")
    template.add_paragraph("{proyecto}", style="Heading 2")
    template.add_paragraph("Preparado por: Empresa Inventada")
    split_field(template.add_paragraph(), ["Para: ", "{", "cli", "ente}", "   Fecha: {fecha}"],
                [None, "bold", "italic", None, None])
    template.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    template.add_paragraph("Índice", style="TOC Heading")
    add_toc(template, ["Cliente", "RESUMEN EJECUTIVO", "Caso 1"], levels=levels, sdt=sdt, dirty=dirty)
    template.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    template.add_heading(MODEL[0], 1)
    template.add_paragraph(MODEL[1])
    template.add_heading(MODEL[2], 2)
    template.add_paragraph(MODEL[3])
    if update_fields:
        flag = OxmlElement("w:updateFields")
        flag.set(qn("w:val"), "true")
        template.settings.element.append(flag)
    template.save(str(path))
    if path.suffix.lower() == ".dotx":
        as_template(path, path)
    return path


def fields_template(path, lines, *, header=None, footer=None):
    template = docx.Document()
    if header:
        template.sections[0].header.paragraphs[0].text = header
    if footer:
        template.sections[0].footer.paragraphs[0].text = footer
    for line in lines:
        template.add_paragraph(line)
    template.save(str(path))
    return path


def every_text(path):
    report = docx.Document(str(path))
    return [layout.paragraph_text(p) for p in report.element.body.iter(qn("w:p"))]


class CoverFieldsTest(Workspace):
    """WI16-AC01 and AC02."""

    LINES = ["Cliente: {cliente}", "Proyecto: {proyecto}", "Reunión: {reunion}", "Fecha: {fecha}", "Tipo: {tipo}",
             "{client} / {project} / {meeting} / {date} / {type} / {reunión} / { Cliente }"]

    def build_with(self, template, **data):
        document.set_template(template, self.data)
        return self.build(**{**MEETING, **data})

    def test_the_five_fields_are_filled_in_spanish_and_in_english(self):
        template = fields_template(self.tmp / "campos.docx", self.LINES)
        result = self.build_with(template)
        self.assertEqual(result.fields, ["client", "project", "meeting", "date", "type"])
        self.assertEqual(self.texts(self.output())[:6], [
            "Cliente: ACME", "Proyecto: Planta Norte", "Reunión: Sesión de dudas", "Fecha: 22 de septiembre de 2026",
            "Tipo: Relevamiento",
            "ACME / Planta Norte / Sesión de dudas / 22 de septiembre de 2026 / Relevamiento / Sesión de dudas / ACME"])
        self.write_summary(summary_text("en"))
        self.build_with(template)
        texts = self.texts(self.output())
        self.assertEqual(texts[3:5], ["Fecha: September 22, 2026", "Tipo: Requirements"])

    def test_fields_in_the_headers_and_footers_are_filled(self):
        template = fields_template(self.tmp / "encabezado.docx", ["Portada"], header="{proyecto} · {cliente}",
                                   footer="{fecha}")
        made = docx.Document(str(template))
        made.sections[0].different_first_page_header_footer = True
        made.sections[0].first_page_header.paragraphs[0].text = "Primera: {tipo}"
        made.save(str(template))
        self.build_with(template)
        section = docx.Document(str(self.output())).sections[0]
        self.assertEqual(section.header.paragraphs[0].text, "Planta Norte · ACME")
        self.assertEqual(section.footer.paragraphs[0].text, "22 de septiembre de 2026")
        self.assertEqual(section.first_page_header.paragraphs[0].text, "Primera: Relevamiento")

    def test_a_meeting_with_no_project_leaves_client_and_project_empty(self):
        template = fields_template(self.tmp / "campos.docx", self.LINES[:2])
        self.build_with(template, project_name=None, client=None)
        self.assertEqual(self.texts(self.output())[:2], ["Cliente: ", "Proyecto: "])

    def test_a_field_keeps_the_templates_format(self):
        template = docx.Document()
        run = template.add_paragraph("Para ").add_run("{cliente}")
        run.bold, run.font.color.rgb, run.font.size = True, RGBColor(0x20, 0x40, 0xA0), Pt(20)
        template.save(str(self.tmp / "formato.docx"))
        self.build_with(self.tmp / "formato.docx")
        filled = docx.Document(str(self.output())).paragraphs[0].runs[1]
        self.assertEqual(filled.text, "ACME")
        self.assertTrue(filled.bold)
        self.assertEqual((filled.font.color.rgb, filled.font.size), (RGBColor(0x20, 0x40, 0xA0), Pt(20)))

    def test_a_field_split_across_runs_is_filled_and_the_rest_stays(self):
        template = docx.Document()
        split_field(template.add_paragraph(), ["Para: ", "{", "cli", "ente}", " fin"],
                    [None, "bold", "italic", None, "italic"])
        split_field(template.add_paragraph(), ["{fe", "cha} y {ti", "po}"], ["bold", None, "italic"])
        template.save(str(self.tmp / "partido.docx"))
        self.build_with(self.tmp / "partido.docx")
        report = docx.Document(str(self.output()))
        first, second = report.paragraphs[0], report.paragraphs[1]
        self.assertEqual(first.text, "Para: ACME fin")
        self.assertEqual([r.text for r in first.runs if r.text], ["Para: ", "ACME", " fin"])
        self.assertTrue(first.runs[1].bold)
        self.assertTrue(first.runs[-1].italic)
        self.assertEqual(second.text, "22 de septiembre de 2026 y Relevamiento")


class FieldRefusalTest(Workspace):
    """WI16-AC03."""

    def test_an_unknown_field_is_refused_when_set_naming_it_and_the_known_ones(self):
        good = fields_template(self.tmp / "buena.docx", ["{cliente}"])
        document.set_template(good, self.data)
        kept = digest(self.data / document.TEMPLATE_NAME)
        with self.assertRaisesRegex(document.ReportError, r"\{clinte\}.*\{cliente\} \{proyecto\}"):
            document.set_template(fields_template(self.tmp / "mala.docx", ["Para: {clinte}"]), self.data)
        self.assertEqual(digest(self.data / document.TEMPLATE_NAME), kept)
        with self.assertRaisesRegex(document.ReportError, r"\{empresa\}"):
            document.set_template(fields_template(self.tmp / "pie.docx", ["x"], footer="{empresa}"), self.data)

    def test_the_report_start_must_be_alone_on_a_line_of_the_body(self):
        for name, lines, footer in (("junto.docx", ["Portada", "Acá {informe} empieza"], None),
                                    ("pie.docx", ["Portada"], "{informe}")):
            with self.subTest(template=name), self.assertRaisesRegex(document.ReportError, "alone on its line"):
                document.set_template(fields_template(self.tmp / name, lines, footer=footer), self.data)

    def test_a_field_left_in_the_finished_document_is_not_delivered(self):
        document.set_template(owner_shaped(self.tmp / "duena.docx"), self.data)
        with mock.patch.object(document.layout, "fill_fields"), \
                self.assertRaisesRegex(document.ReportError, r"\{cliente\}, \{fecha\}, \{proyecto\} unfilled"):
            self.build(**MEETING)
        self.assertNothingWritten()

    def test_a_template_placed_by_hand_with_an_unknown_field_is_refused_when_building(self):
        self.data.mkdir()
        fields_template(self.data / document.TEMPLATE_NAME, ["{clinte}"])
        with self.assertRaisesRegex(document.ReportError, r"\{clinte\}"):
            self.build(**MEETING)
        self.assertNothingWritten()


class ReportStartTest(Workspace):
    """WI16-AC04."""

    def test_the_model_after_the_table_of_contents_is_not_in_the_report(self):
        document.set_template(owner_shaped(self.tmp / "duena.docx"), self.data)
        result = self.build(**MEETING)
        texts = every_text(self.output())
        self.assertEqual(texts[:4], ["ACME", "Planta Norte", "Preparado por: Empresa Inventada",
                                     "Para: ACME   Fecha: 22 de septiembre de 2026"])
        self.assertIn("Índice", texts)
        for model in MODEL + ["Cliente", "Caso 1"]:
            self.assertNotIn(model, texts)
        self.assertEqual(result.dropped, 4)
        self.assertEqual(result.contents, len(document.summary_headings(summary_text())))
        self.assertIn(MEETING["title"], texts)

    def test_the_model_after_informe_is_not_in_the_report(self):
        template = fields_template(self.tmp / "marca.docx", ["Portada de {cliente}", "{informe}", "Modelo que no va"])
        made = docx.Document(str(template))
        made.add_heading("Un título de modelo", 1)
        made.save(str(template))
        document.set_template(template, self.data)
        result = self.build(**MEETING)
        texts = every_text(self.output())
        self.assertEqual(texts[0], "Portada de ACME")
        for gone in ("{informe}", "Modelo que no va", "Un título de modelo"):
            self.assertNotIn(gone, texts)
        self.assertEqual((result.dropped, result.contents), (2, 0))

    def test_a_table_of_contents_inside_a_content_control_is_found(self):
        document.set_template(owner_shaped(self.tmp / "control.docx", sdt=True), self.data)
        result = self.build(**MEETING)
        texts = every_text(self.output())
        for model in MODEL:
            self.assertNotIn(model, texts)
        report = docx.Document(str(self.output()))
        (toc,) = layout._tocs(report)
        self.assertEqual(toc[0].getparent().tag, qn("w:sdtContent"))
        self.assertEqual(layout.toc_entries(report, toc), document.summary_headings(summary_text()))
        self.assertEqual(result.dropped, 4)


class TableOfContentsTest(Workspace):
    """WI16-AC05 and AC06."""

    def build_owner_shaped(self, **options):
        document.set_template(owner_shaped(self.tmp / "duena.docx", **options), self.data)
        return self.build(**MEETING)

    def start_of(self, report):
        children = layout.body_children(report)
        return next(number for number, child in enumerate(children)
                    if layout.paragraph_text(child) == MEETING["title"])

    def test_the_table_lists_the_reports_headings_each_linking_to_its_heading(self):
        self.build_owner_shaped()
        report = docx.Document(str(self.output()))
        (toc,) = layout._tocs(report)
        headings = document.summary_headings(summary_text())
        self.assertEqual(layout.toc_entries(report, toc), headings)
        self.assertIn('TOC \\o "1-3"', toc[2])
        bookmarks = {node.get(qn("w:name")): layout._paragraph_of(node)
                     for node in report.element.body.iter(qn("w:bookmarkStart"))}
        links = [link for p in layout.paragraphs([toc[0].getparent()]) for link in p.iter(qn("w:hyperlink"))]
        self.assertEqual(len(links), len(headings))
        for link, heading in zip(links, headings):
            target = bookmarks[link.get(qn("w:anchor"))]
            self.assertEqual(layout.paragraph_text(target), heading)
            self.assertTrue(docx.text.paragraph.Paragraph(target, report.part).style.name.startswith("Heading"))
        kinds = [node.get(qn("w:fldCharType")) for node in report.element.body.iter(qn("w:fldChar"))]
        self.assertEqual(kinds, ["begin", "separate", "end"])

    def test_the_levels_the_field_asks_for_are_kept(self):
        self.write_summary(summary_text() + "\n### Detalle del costo\nTexto del detalle.\n")
        headings = document.summary_headings(summary_text())
        for levels, expected in (("1-1", headings), ("1-3", headings + ["Detalle del costo"])):
            with self.subTest(levels=levels):
                self.build_owner_shaped(levels=levels)
                report = docx.Document(str(self.output()))
                self.assertEqual(layout.toc_entries(report, layout._tocs(report)[0]), expected)

    def test_word_is_not_asked_to_update_fields_on_opening(self):
        self.build_owner_shaped(dirty=True, update_fields=True)
        report = docx.Document(str(self.output()))
        self.assertEqual(report.settings.element.findall(qn("w:updateFields")), [])
        self.assertEqual([node for node in report.element.body.iter(qn("w:fldChar")) if node.get(qn("w:dirty"))], [])

    def check(self, path, start):
        headings = document.summary_headings(summary_text())
        frames = [self.frames / name for name in list(FRAMES)[:2]]
        document.check_report(path, headings, [], frames, start=start, contents=[headings])

    def test_a_section_missing_from_the_body_is_not_delivered_even_if_the_table_lists_it(self):
        self.build_owner_shaped()
        report = docx.Document(str(self.output()))
        start = self.start_of(report)
        self.check(self.output(), start)
        heading = next(child for child in layout.body_children(report)[start:]
                       if layout.paragraph_text(child) == writer.required_headings("es")[2])
        heading.getparent().remove(heading)
        report.save(str(self.tmp / "sin-seccion.docx"))
        with self.assertRaisesRegex(document.ReportError, "missing the section"):
            self.check(self.tmp / "sin-seccion.docx", start)

    def test_a_table_of_contents_that_does_not_list_exactly_the_sections_is_not_delivered(self):
        self.build_owner_shaped()
        report = docx.Document(str(self.output()))
        start = self.start_of(report)
        (toc,) = layout._tocs(report)
        entry = toc[0].getnext()
        entry.getparent().remove(entry)
        report.save(str(self.tmp / "indice-corto.docx"))
        with self.assertRaisesRegex(document.ReportError, "table of contents"):
            self.check(self.tmp / "indice-corto.docx", start)


class TemplateInfoTest(Workspace):
    """WI16-AC07."""

    run_main = CommandLineTest.run_main

    def test_the_real_name_and_what_was_understood_are_kept(self):
        document.set_template(owner_shaped(self.tmp / "Plantilla G7.dotx"), self.data)
        info = document.template_info(self.data)
        self.assertEqual(info.name, "Plantilla G7.dotx")
        self.assertRegex(info.set_utc, r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$")
        self.assertEqual((info.fields, info.tables_of_contents, info.start, info.dropped),
                         (["client", "project", "date"], 1, "index", 4))
        code, out, _ = self.run_main("template", "show")
        self.assertEqual(code, 0)
        for expected in ("Plantilla G7.dotx", "fills: client, project, date", "table of contents: yes",
                         "dropped after its table of contents: 4 paragraph(s)"):
            self.assertIn(expected, out)

    def test_the_name_given_is_the_one_kept(self):
        upload = owner_shaped(self.tmp / "a1b2c3.docx")
        document.set_template(upload, self.data, name="C:\\Users\\x\\Mi plantilla.docx")
        self.assertEqual(document.template_info(self.data).name, "Mi plantilla.docx")

    def test_a_template_kept_before_this_work_item_says_its_name_was_not_recorded(self):
        self.data.mkdir()
        company_template(self.data / document.TEMPLATE_NAME)
        info = document.template_info(self.data)
        self.assertEqual((info.name, info.set_utc, info.start, info.dropped), (None, None, None, 0))
        code, out, _ = self.run_main("template", "show")
        self.assertIn("its original name was not recorded", out)
        self.assertIn("cover: everything on its page", out)

    def test_remove_removes_the_record_too(self):
        document.set_template(owner_shaped(self.tmp / "duena.docx"), self.data)
        self.assertTrue((self.data / document.TEMPLATE_RECORD).is_file())
        document.remove_template(self.data)
        self.assertEqual(list(self.data.iterdir()), [])


class ExampleTemplateTest(Workspace):
    """WI16-AC08."""

    run_main = CommandLineTest.run_main

    def test_the_example_can_be_set_and_fills_every_field(self):
        example = self.tmp / "ejemplo.docx"
        code, out, _ = self.run_main("template", "example", str(example))
        self.assertEqual(code, 0, out)
        code, out, _ = self.run_main("template", "set", str(example))
        self.assertEqual(code, 0, out)
        result = self.build(**MEETING)
        self.assertEqual(sorted(result.fields), ["client", "date", "meeting", "project", "type"])
        self.assertEqual(result.contents, len(document.summary_headings(summary_text())))
        texts = every_text(self.output())
        self.assertEqual(texts[:5], ["Planta Norte", "Cliente: ACME", "Reunión: Sesión de dudas",
                                     "Tipo de reunión: Relevamiento", "Fecha: 22 de septiembre de 2026"])
        self.assertFalse([text for text in texts if "Cómo usar esta plantilla" in text or "modelo" in text])

    def test_template_example_does_not_write_over_a_file(self):
        example = self.tmp / "ejemplo.docx"
        example.write_bytes(b"mine")
        code, _, err = self.run_main("template", "example", str(example))
        self.assertEqual((code, example.read_bytes()), (2, b"mine"))
        self.assertIn("already exists", err)


if __name__ == "__main__":
    unittest.main()
