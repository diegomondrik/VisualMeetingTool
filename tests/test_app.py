"""Tests for the application (meetingtool.app, WI15). The server runs on a free
port of 127.0.0.1 over a data folder invented at test time; Gemini is the fake
on localhost from test_reading, and the key lives in a dictionary instead of
the Windows Credential Manager: no test reaches the network or needs a key."""

import contextlib
import http.client
import io
import json
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest import mock
from urllib.parse import quote

import docx
import numpy as np
from PIL import Image

from meetingtool.app import company, jobs, library, pages, server
from meetingtool.projects import store
from meetingtool.reading import gemini
from meetingtool.report import document
from meetingtool.summary import qa, writer
from meetingtool.summary.__main__ import main as summary_main
from tests import test_qa, test_report, test_summary
from tests.test_frames import SLIDES, write_teams_docx, write_video
from tests.test_reading import KEY, FakeGemini, answer_for

REPOSITORY = Path(__file__).resolve().parent.parent
JSON = {"Content-Type": "application/json", "X-MeetingTool": "1"}
RAW = {"Content-Type": "application/octet-stream", "X-MeetingTool": "1"}
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FRAME = "frame_007_t00-12-34.jpg"
SUMMARY_MD = ("## Resumen ejecutivo\n\nSe revisó el costo de proceso <script>alert(1)</script>.\n\n"
              f"## Qué se mostró\n\n- La planilla de tarifas [{FRAME}]\n\n| Etapa | Tarifa |\n|---|---|\n| Faena | 3 |\n")


class Keys:
    """The Windows Credential Manager, in a dictionary."""

    def __init__(self, key=None):
        self.key = key

    def read(self):
        return self.key

    def save(self, key):
        self.key = key

    def delete(self):
        had, self.key = self.key is not None, None
        return had


class Running(unittest.TestCase):
    """A running application over a data folder of its own."""

    key = KEY

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.data = self.tmp / "data"
        self.data.mkdir()
        self.keys = Keys(self.key)
        self.opened = []
        self.fake = FakeGemini(self.script()).__enter__()
        self.app = server.App(self.data, read_key=self.keys.read, save_key=self.keys.save, delete_key=self.keys.delete,
                              endpoint=self.fake.endpoint, sleep=lambda seconds: None, retry_delays=(0, 0),
                              opener=self.opened.append).start()

    def tearDown(self):
        self.app.stop()
        self.fake.__exit__(None, None, None)
        self._tmp.cleanup()

    def script(self):
        return ()

    def request(self, method, path, body=None, headers=None, cookie=True, host=None):
        connection = http.client.HTTPConnection(server.HOST, self.app.port, timeout=60)
        sent = dict(headers or {})
        if cookie is True:
            sent["Cookie"] = f"{server.COOKIE}={self.app.token}"
        elif cookie:
            sent["Cookie"] = cookie
        if host is not None:
            sent["Host"] = host
        if isinstance(body, (dict, list)):
            body = json.dumps(body)
        connection.request(method, path, body, sent)
        response = connection.getresponse()
        data = response.read()
        connection.close()
        return response.status, {k.lower(): v for k, v in response.getheaders()}, data

    def page(self, path):
        status, headers, data = self.request("GET", path)
        self.assertEqual(status, 200, data[:300])
        return data.decode("utf-8")

    def api(self, path, data=None, expect=200):
        status, _, body = self.request("POST", path, data or {}, JSON)
        self.assertEqual(status, expect, body[:300])
        return json.loads(body)

    def upload(self, kind, path, name=None):
        status, _, body = self.request("PUT", f"/api/upload?kind={kind}&name={quote(name or path.name)}",
                                       path.read_bytes(), RAW)
        self.assertEqual(status, 200, body)
        return json.loads(body)["upload"]

    def new_project(self, name="Cermaq Sprint 3", client="Cermaq"):
        return self.api("/api/projects", {"name": name, "client": client})["id"]


# ── WI15-AC01: the five screens ───────────────────────────────────────────────

class ScreensTest(Running):
    def setUp(self):
        super().setUp()
        self.project = self.new_project()
        folder = self.data / self.project / library.RESULTS_DIR / "2026-09-25-dudas-a1b2c3"
        folder.mkdir(parents=True)
        (folder / writer.OUTPUT_NAME).write_text(SUMMARY_MD, encoding="utf-8")
        Image.new("RGB", (64, 36), (20, 90, 200)).save(folder / FRAME)
        Image.new("RGB", (64, 36), (200, 90, 20)).save(folder / "frame_001_t00-00-10.jpg")
        (folder / library.REPORT_NAME).write_bytes(b"PK\x03\x04 a report")
        (folder / library.RUN_NAME).write_text(json.dumps({
            "format": "summary", "cost_usd": 0.1234, "max_cost_usd": 0.5, "seconds": 321.0,
            "stages": [{"name": "frames", "label": "Imágenes del video", "state": "done", "seconds": 300.0,
                        "cost_usd": 0.0},
                       {"name": "summary", "label": "Resumen", "state": "done", "seconds": 21.0,
                        "cost_usd": 0.1234}]}), encoding="utf-8")
        self.processed = store.add_meeting(self.data, self.project, "Sesión de dudas", "2026-09-25", "requirements",
                                           summary="Se revisaron las dudas.", key_points=["Juan manda el detalle"],
                                           meeting_folder=f"{library.RESULTS_DIR}/{folder.name}")
        self.terminal = store.add_meeting(self.data, self.project, "Relevamiento inicial", "2026-09-10", "presale",
                                          summary="El primer relevamiento.", key_points=["Hay tres plantas"])
        loose = self.data / "d181-relevamiento"
        loose.mkdir()
        (loose / writer.OUTPUT_NAME).write_text(f"## Resumen\n\nUna reunión suelta [{FRAME}].\n", encoding="utf-8")
        Image.new("RGB", (64, 36), (10, 10, 10)).save(loose / FRAME)
        (self.data / "reunion2-cermaq").mkdir()  # no summary: not a result
        (self.data / "reunion2-cermaq" / "transcript.txt").write_text("[00:00:01] hola", encoding="utf-8")

    def test_the_projects_screen_lists_projects_and_loose_results(self):
        text = self.page("/")
        self.assertIn('href="/p/cermaq-sprint-3"', text)
        self.assertIn("Cermaq Sprint 3", text)
        self.assertIn("<td>Cermaq</td>", text)
        self.assertIn('href="/r/d181-relevamiento"', text)
        self.assertNotIn("reunion2-cermaq", text)
        self.assertIn('data-api="/api/projects"', text)

    def test_a_project_is_created_from_the_page(self):
        created = self.new_project("Proyecto Nuevo", "Otra empresa")
        self.assertEqual(created, "proyecto-nuevo")
        self.assertIn("Proyecto Nuevo", self.page("/"))
        self.assertIn("Otra empresa", self.page("/p/proyecto-nuevo"))
        self.api("/api/projects", {"name": "Proyecto Nuevo", "client": ""}, expect=400)

    def test_the_project_screen_lists_meetings_in_date_order_with_their_knowledge(self):
        text = self.page(f"/p/{self.project}")
        first, second = text.index("Relevamiento inicial"), text.index("Sesión de dudas")
        self.assertLess(first, second)
        self.assertIn("Relevamiento", text)
        self.assertIn("Preventa", text)
        self.assertIn("US$0,123", text)
        self.assertIn(f'href="/p/{self.project}/new"', text)
        knowledge = text[text.index("Lo que el proyecto ya sabe"):]
        self.assertIn("El primer relevamiento.", knowledge)
        self.assertIn("Juan manda el detalle", knowledge)

    def test_a_processed_meeting_shows_its_summary_frames_report_and_cost(self):
        text = self.page(f"/p/{self.project}/m/{self.processed['id']}")
        self.assertIn("<h3>Resumen ejecutivo</h3>", text)  # the page's own title is the only h1
        self.assertIn(f'src="/p/{self.project}/m/{self.processed["id"]}/f/{FRAME}"', text)
        self.assertIn("Minuto 12:34", text)
        self.assertNotIn("frame_001_t00-00-10.jpg", text)  # kept, but not named by the summary
        self.assertIn("1 imágenes en el informe, de 2", text)
        self.assertIn(f'data-open="{self.project}/{self.processed["id"]}"', text)
        self.assertIn("summary.docx\" download", text)
        self.assertIn("US$0,123", text)
        self.assertIn("<th>Etapa</th><th>Tarifa</th>", text)

    def test_what_the_summary_says_is_escaped(self):
        text = self.page(f"/p/{self.project}/m/{self.processed['id']}")
        self.assertNotIn("<script>alert", text)
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", text)
        self.assertEqual(re.findall(r"<script(?![^>]*\bsrc=)", text), [])

    def test_a_meeting_added_from_the_terminal_shows_what_the_project_kept(self):
        text = self.page(f"/p/{self.project}/m/{self.terminal['id']}")
        self.assertIn("El primer relevamiento.", text)
        self.assertIn("Hay tres plantas", text)
        self.assertIn("se cargó desde la terminal", text)
        self.assertNotIn("data-open", text)

    def test_the_files_of_a_meeting_are_served_and_nothing_else(self):
        base = f"/p/{self.project}/m/{self.processed['id']}/f/"
        status, headers, data = self.request("GET", base + FRAME)
        self.assertEqual((status, headers["content-type"]), (200, "image/jpeg"))
        folder = self.data / self.project / self.processed["folder"]
        self.assertEqual(data, (folder / FRAME).read_bytes())
        status, headers, data = self.request("GET", base + library.REPORT_NAME)
        self.assertEqual((status, headers["content-type"]), (200, DOCX_TYPE))
        self.assertIn("attachment", headers["content-disposition"])
        for name in ("run.json", writer.OUTPUT_NAME, "..%2F..%2Fproject.json", "%2E%2E", "frame_7.jpg"):
            self.assertEqual(self.request("GET", base + name)[0], 404, name)
        self.assertEqual(self.request("GET", f"/p/{self.project}/m/{self.terminal['id']}/f/{FRAME}")[0], 404)
        self.assertEqual(self.request("GET", "/p/..%2F..%2Fetc/m/x")[0], 404)

    def test_a_record_naming_a_folder_outside_its_results_is_not_followed(self):
        for folder in ("../d181-relevamiento", "results/../../d181-relevamiento", "C:/Windows", "meetings"):
            record = store.add_meeting(self.data, self.project, f"Falsa {len(folder)}", "2026-09-26",
                                       meeting_folder=folder)
            text = self.page(f"/p/{self.project}/m/{record['id']}")
            self.assertNotIn("Una reunión suelta", text, folder)
            self.assertEqual(self.request("GET", f"/p/{self.project}/m/{record['id']}/f/{FRAME}")[0], 404, folder)

    def test_a_loose_result_is_shown_read_only(self):
        text = self.page("/r/d181-relevamiento")
        self.assertIn("Una reunión suelta", text)
        self.assertIn(f'src="/r/d181-relevamiento/f/{FRAME}"', text)
        self.assertNotIn("data-open", text)  # it has no report
        self.assertEqual(self.request("GET", f"/r/d181-relevamiento/f/{FRAME}")[0], 200)
        self.assertEqual(self.request("GET", "/r/reunion2-cermaq")[0], 404)
        self.assertEqual(self.request("GET", "/r/reunion2-cermaq/f/transcript.txt")[0], 404)

    def test_the_word_report_is_opened_on_the_machine(self):
        self.api("/api/open", {"target": f"{self.project}/{self.processed['id']}"})
        self.assertEqual(self.opened, [self.data / self.project / self.processed["folder"] / library.REPORT_NAME])
        self.api("/api/open", {"target": f"{self.project}/{self.terminal['id']}"}, expect=404)
        self.api("/api/open", {"target": "d181-relevamiento"}, expect=404)
        self.api("/api/open", {"target": "../x/y"}, expect=404)
        self.assertEqual(len(self.opened), 1)

    def test_the_new_meeting_screen_has_every_field(self):
        text = self.page(f"/p/{self.project}/new")
        for field in ('name="title"', 'name="date"', 'name="transcript"', 'name="recording"', 'name="meeting_type"',
                      'name="language"', 'value="summary"', 'value="qa"', 'name="max_cost"', 'value="1.00"'):
            self.assertIn(field, text)
        for meeting_type in writer.MEETING_TYPES:
            self.assertIn(f'value="{meeting_type}"', text)
        self.assertIn('data-project="cermaq-sprint-3"', text)
        self.assertEqual(self.request("GET", "/p/no-existe/new")[0], 404)


class SettingsTest(Running):
    key = None

    def test_the_key_is_saved_replaced_and_deleted_and_never_shown(self):
        self.assertIn("No hay una clave guardada", self.page("/settings"))
        self.api("/api/key", {"key": f"  {KEY}  "})
        self.assertEqual(self.keys.key, KEY)
        text = self.page("/settings")
        self.assertIn(f"Hay una clave guardada ({len(KEY)} caracteres)", text)
        self.assertNotIn(KEY, text)
        self.assertNotIn("k3yk3y", text)
        self.api("/api/key", {"key": "AIzaOTRA"})
        self.assertEqual(self.keys.key, "AIzaOTRA")
        self.api("/api/key/delete")
        self.assertIsNone(self.keys.key)
        self.assertIn("No hay una clave guardada", self.page("/settings"))

    def test_a_key_that_cannot_be_a_key_is_not_saved(self):
        self.api("/api/key", {"key": "clave con ñ"}, expect=400)
        self.api("/api/key", {"key": "   "}, expect=400)
        self.api("/api/key", {"key": 12}, expect=400)
        self.assertIsNone(self.keys.key)

    def test_the_company_template_is_set_and_removed(self):
        self.assertIn("Sin plantilla", self.page("/settings"))
        template = self.tmp / "empresa.dotx"
        made = docx.Document()
        made.add_paragraph("Empresa Inventada S.A.")
        made.save(self.tmp / "empresa.docx")
        shutil.copy(self.tmp / "empresa.docx", template)
        status, _, body = self.request("PUT", "/api/template?name=empresa.docx", (self.tmp / "empresa.docx").read_bytes(),
                                       RAW)
        self.assertEqual(status, 200, body)
        self.assertEqual(document.stored_template(self.data), self.data / document.TEMPLATE_NAME)
        self.assertIn("empresa.docx", self.page("/settings"))  # WI16-AC07: the file's name, not the copy's
        self.assertEqual(list((self.data / ".meetingtool-uploads").iterdir()), [])
        status, _, body = self.request("PUT", "/api/template?name=macros.docm", b"PK", RAW)
        self.assertEqual(status, 415)
        status, _, body = self.request("PUT", "/api/template?name=roto.docx", b"not a word file", RAW)
        self.assertEqual(status, 400, body)
        self.assertEqual(list((self.data / ".meetingtool-uploads").iterdir()), [])
        self.api("/api/template/remove")
        self.assertIsNone(document.stored_template(self.data))

    def test_the_settings_say_the_templates_real_name_and_what_was_understood(self):
        """WI16-AC07."""
        template = test_report.owner_shaped(self.tmp / "Plantilla G7.dotx")
        status, _, body = self.request("PUT", "/api/template?name=" + quote("Plantilla G7.dotx"),
                                       template.read_bytes(), RAW)
        self.assertEqual((status, json.loads(body)), (200, {"template": "Plantilla G7.dotx"}))
        page = self.page("/settings")
        self.assertIn("<strong>Plantilla G7.dotx</strong>, cargada el ", page)
        self.assertNotIn(document.TEMPLATE_NAME, page)
        for said in ("En la portada va a poner: el cliente, el proyecto, la fecha.", "Tiene índice",
                     "Lo que tiene después del índice es un modelo y no entra en los informes: 4 párrafos con texto."):
            self.assertIn(said, page)
        status, _, body = self.request("PUT", "/api/template?name=mala.docx", test_report.fields_template(
            self.tmp / "mala.docx", ["{clinte}"]).read_bytes(), RAW)
        self.assertEqual(status, 400)
        self.assertIn("{clinte}", json.loads(body)["error"])
        self.assertIn("Plantilla G7.dotx", self.page("/settings"))

    def test_a_kept_template_that_can_no_longer_be_used_is_said_and_can_be_removed(self):
        (self.data / document.TEMPLATE_NAME).write_bytes(b"not a word file")
        page = self.page("/settings")
        self.assertIn("La plantilla guardada ya no se puede usar", page)
        self.assertIn("Dejar de usar la plantilla", page)
        self.api("/api/template/remove")
        self.assertIn("Sin plantilla", self.page("/settings"))

    def test_the_example_template_downloads_and_can_be_set(self):
        """WI16-AC08."""
        self.assertIn(f'href="/{pages.EXAMPLE_NAME}" download', self.page("/settings"))
        status, headers, body = self.request("GET", "/" + pages.EXAMPLE_NAME)
        self.assertEqual((status, headers["content-type"]), (200, DOCX_TYPE))
        self.assertIn(pages.EXAMPLE_NAME, headers["content-disposition"])
        status, _, _ = self.request("GET", "/" + pages.EXAMPLE_NAME, cookie=None)
        self.assertEqual(status, 403)
        status, _, answer = self.request("PUT", "/api/template?name=" + pages.EXAMPLE_NAME, body, RAW)
        self.assertEqual(status, 200, answer)
        self.assertEqual(document.template_info(self.data).fields, ["project", "client", "meeting", "type", "date"])


# ── WI17-AC01 and AC02: the company's name and logo ───────────────────────────

def png_bytes(size=(40, 20), color=(200, 10, 10)):
    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, "PNG")
    return buffer.getvalue()


class CompanyTest(Running):
    SCREENS = ("/", "/settings", "/p/{project}", "/p/{project}/new", "/no-such-page")

    def setUp(self):
        super().setUp()
        self.project = self.new_project()

    def logo(self, data, name="logo.png"):
        return self.request("PUT", "/api/logo?name=" + quote(name), data, RAW)

    def screens(self):
        return {path: self.request("GET", path.format(project=self.project))[2].decode("utf-8")
                for path in self.SCREENS}

    def assertNoLogo(self):
        self.assertIsNone(company.logo_path(self.data))
        self.assertEqual(self.request("GET", "/company/logo")[0], 404)
        uploads = self.data / ".meetingtool-uploads"
        self.assertEqual(list(uploads.iterdir()) if uploads.exists() else [], [])

    def test_the_name_and_the_logo_are_on_every_screen_and_in_the_tabs_name(self):
        self.api("/api/company", {"name": "  Nexo   Consultores  "})
        status, _, body = self.logo(png_bytes())
        self.assertEqual(status, 200, body)
        for path, page in self.screens().items():
            self.assertIn('<a class="brand" href="/"><img class="logo" src="/company/logo" alt="">'
                          '<span class="company">Nexo Consultores</span><span class="product">MeetingTool</span></a>',
                          page, path)
            self.assertRegex(page, r"<title>[^<]* · Nexo Consultores · MeetingTool</title>", path)
        status, headers, data = self.request("GET", "/company/logo")
        self.assertEqual((status, headers["content-type"]), (200, "image/png"))
        self.assertEqual(Image.open(io.BytesIO(data)).size, (40, 20))
        self.assertEqual(self.request("GET", "/company/logo", cookie=False)[0], 403)

    def test_the_name_and_the_logo_change_and_go(self):
        self.api("/api/company", {"name": "Nexo"})
        self.logo(png_bytes())
        self.api("/api/company", {"name": "Otra"})
        jpeg = io.BytesIO()
        Image.new("RGB", (30, 30), (0, 90, 20)).save(jpeg, "JPEG")
        status, _, body = self.logo(jpeg.getvalue(), "nuevo.JPG")
        self.assertEqual(status, 200, body)
        self.assertEqual(company.logo_path(self.data).name, "company-logo.jpg")
        self.assertFalse((self.data / "company-logo.png").exists())
        self.assertEqual(self.request("GET", "/company/logo")[1]["content-type"], "image/jpeg")
        self.assertIn('<span class="company">Otra</span>', self.page("/"))
        self.assertEqual(self.api("/api/logo/remove"), {"removed": True})
        self.api("/api/company", {"name": ""})
        self.assertNoLogo()
        self.assertIn('<a class="brand" href="/">MeetingTool</a>', self.page("/"))

    def test_without_a_name_or_a_logo_the_screens_are_as_before(self):
        before = self.screens()
        self.api("/api/company", {"name": "Nexo"})
        self.logo(png_bytes())
        self.api("/api/logo/remove")
        self.api("/api/company", {"name": "   "})
        self.assertEqual(self.screens(), before)
        self.assertIn('<a class="brand" href="/">MeetingTool</a>', before["/"])
        self.assertIn("<title>Proyectos · MeetingTool</title>", before["/"])

    def test_a_name_with_markup_is_shown_as_text(self):
        self.api("/api/company", {"name": '<script>alert(1)</script> & "x"'})
        for path, page in self.screens().items():
            self.assertNotIn("<script>alert", page, path)
            self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt; &amp; &quot;x&quot;", page, path)
            self.assertEqual(re.findall(r"<script(?![^>]*\bsrc=)", page), [], path)

    def test_a_name_that_cannot_be_shown_is_refused(self):
        for name, said in (("x" * 81, "hasta 80 caracteres"), ("Nexo\x07", "no se pueden mostrar"),
                           ("Ne‮xo", "no se pueden mostrar"), (12, "es texto")):
            self.assertIn(said, self.api("/api/company", {"name": name}, expect=400)["error"])
        self.assertEqual(company.company(self.data).name, "")

    def test_a_logo_that_is_not_a_real_png_or_jpg_is_refused_and_nothing_is_kept(self):
        whole = png_bytes((200, 120))
        gif = io.BytesIO()
        Image.new("RGB", (10, 10)).save(gif, "GIF")
        svg = b'<?xml version="1.0"?><svg xmlns="http://www.w3.org/2000/svg" onload="alert(1)"/>'
        for data, name, status, said in (
                (b"just some text, not an image", "logo.png", 400, "no es una imagen PNG o JPG"),
                (whole[:len(whole) // 2], "logo.png", 400, "no es una imagen PNG o JPG"),
                (svg, "logo.svg", 415, "SVG no se acepta"),
                (svg, "logo.png", 400, "SVG no se acepta"),
                (b"  <svg><script>alert(1)</script></svg>", "logo.jpg", 400, "SVG no se acepta"),
                (gif.getvalue(), "logo.png", 400, "tiene que ser un PNG o un JPG"),
                (gif.getvalue(), "logo.gif", 415, "se acepta .png, .jpg, .jpeg"),
                (png_bytes((4001, 10)), "logo.png", 400, "hasta 4000 píxeles")):
            answer_status, _, body = self.logo(data, name)
            self.assertEqual(answer_status, status, (name, body))
            self.assertIn(said, json.loads(body)["error"], name)
            self.assertNoLogo()

    def test_a_logo_of_more_than_1_mb_is_refused_saying_why(self):
        noise = Image.fromarray(np.random.default_rng(7).integers(0, 255, (700, 700, 3), dtype=np.uint8))
        buffer = io.BytesIO()
        noise.save(buffer, "PNG")
        self.assertGreater(len(buffer.getvalue()), company.LOGO_LIMIT)
        status, _, body = self.logo(buffer.getvalue())
        self.assertEqual(status, 413)
        self.assertIn("pesa más de 1 MB", json.loads(body)["error"])
        with self.assertRaises(company.SettingsError) as caught:
            company.check_logo(buffer.getvalue(), "logo.png")
        self.assertEqual(caught.exception.message.key, "app.logo.too_big")
        self.assertNoLogo()

    def test_a_logo_of_more_than_1_mb_gets_its_reason_every_time(self):
        """The body of a refused logo is read before the answer, so the page
        gets the reason, not a cut connection (review of 5d4c63a, P1-1: 3 of 8
        runs were cut on Windows)."""
        big = b"\x89PNG" + bytes(company.LOGO_LIMIT + 400_000)
        for _ in range(20):
            status, _, body = self.logo(big)
            self.assertEqual(status, 413)
            self.assertIn("pesa más de 1 MB", json.loads(body)["error"])
        self.assertNoLogo()

    def test_a_refused_request_learns_nothing_of_the_company(self):
        """A request without the session, or for another host, gets the refusal
        in the application's language, without the company's name or logo
        (review of 5d4c63a, P2-1)."""
        self.api("/api/company", {"name": "Nexo Secreto"})
        self.logo(png_bytes())
        for cookie, host in ((False, None), (f"{server.COOKIE}=wrong", None), (True, "evil.example")):
            status, _, body = self.request("GET", "/", cookie=cookie, host=host)
            self.assertEqual(status, 403)
            self.assertNotIn(b"Nexo", body)
            self.assertNotIn(b"/company/logo", body)
        status, _, body = self.request("GET", "/open?token=wrong", cookie=False)
        self.assertEqual(status, 403)
        self.assertNotIn(b"Nexo", body)

    def test_only_the_pixels_of_the_logo_are_kept(self):
        """A colour profile, text chunks or metadata of the file are not kept
        (review of 5d4c63a, P2-4); a transparent colour is part of the image."""
        from PIL import PngImagePlugin

        info = PngImagePlugin.PngInfo()
        info.add_text("Comment", "<script>alert(1)</script>")
        image = Image.new("P", (20, 10), 1)
        image.putpalette([0, 0, 0, 255, 0, 0] + [0] * 762)
        buffer = io.BytesIO()
        image.save(buffer, "PNG", pnginfo=info, icc_profile=b"CHOSEN-BYTES" * 40, transparency=1)
        status, _, body = self.logo(buffer.getvalue())
        self.assertEqual(status, 200, body)
        kept = company.logo_path(self.data).read_bytes()
        for chunk in (b"iCCP", b"tEXt", b"CHOSEN", b"<script>"):
            self.assertNotIn(chunk, kept)
        self.assertEqual(Image.open(io.BytesIO(kept)).info.get("transparency"), 1)
        status, _, body = self.logo(b"just some text, not an image")
        self.assertNotIn("object at 0x", json.dumps(json.loads(body)["detail"]))

    def test_what_is_kept_is_the_image_written_again(self):
        """Anything a file carries besides its image (here, a page appended to a
        PNG) is not kept or served."""
        status, _, body = self.logo(png_bytes() + b"<html><script>alert(1)</script></html>")
        self.assertEqual(status, 200, body)
        kept = company.logo_path(self.data).read_bytes()
        self.assertNotIn(b"<script>", kept)
        self.assertEqual(Image.open(io.BytesIO(kept)).size, (40, 20))


# ── WI15-AC02: only this machine, and no other site ───────────────────────────

class IsolationTest(Running):
    def setUp(self):
        super().setUp()
        self.project = self.new_project()
        store.add_meeting(self.data, self.project, "Sesión de dudas", "2026-09-25", summary="Dato del cliente.")

    def assertNothingCreated(self):
        self.assertEqual([p["id"] for p in store.list_projects(self.data)], [self.project])

    def test_the_server_listens_only_on_the_loopback_address(self):
        self.assertEqual(self.app.httpd.server_address[0], "127.0.0.1")
        others = {info[4][0] for info in socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET)}
        others = {address for address in others if not address.startswith("127.")}
        if not others:
            self.skipTest("this machine has no address other than 127.0.0.1")
        for address in others:
            with self.assertRaises(OSError, msg=address):
                socket.create_connection((address, self.app.port), timeout=3).close()

    def test_a_peer_that_is_not_this_machine_is_refused(self):
        headers = {"Host": f"127.0.0.1:{self.app.port}", "Cookie": f"{server.COOKIE}={self.app.token}"}
        self.assertIsNone(server.refusal(self.app, "127.0.0.1", "GET", "/", headers))
        self.assertIsNone(server.refusal(self.app, "::1", "GET", "/", headers))
        for peer in ("192.168.1.20", "10.0.0.5", "8.8.8.8", "fe80::1", "not an address"):
            self.assertEqual(server.refusal(self.app, peer, "GET", "/", headers)[0], 403, peer)

    def test_without_the_session_cookie_nothing_is_read(self):
        paths = ["/", f"/p/{self.project}", f"/p/{self.project}/new", "/settings", "/api/running",
                 "/static/app.js", "/r/x", "/job/x"]
        for cookie in (False, f"{server.COOKIE}=wrong", f"{server.COOKIE}={self.app.token[:-1]}", "other=1",
                       "\x00bad"):
            for path in paths:
                status, _, data = self.request("GET", path, cookie=cookie)
                self.assertEqual(status, 403, (cookie, path))
                self.assertNotIn(b"Cermaq", data)
                self.assertNotIn(b"Dato del cliente", data)

    def test_without_the_session_cookie_nothing_is_changed(self):
        status, _, _ = self.request("POST", "/api/projects", {"name": "Intruso"}, JSON, cookie=False)
        self.assertEqual(status, 403)
        status, _, _ = self.request("POST", "/api/key", {"key": "AIzaINTRUSO"}, JSON, cookie=False)
        self.assertEqual(status, 403)
        self.assertEqual(self.keys.key, KEY)
        self.assertNothingCreated()

    def test_another_host_name_is_refused_even_with_the_cookie(self):
        # A DNS rebinding: a site's name made to point at 127.0.0.1.
        for host in (f"evil.example:{self.app.port}", "evil.example", f"127.0.0.1:{self.app.port + 1}",
                     f"127.0.0.1.evil.example:{self.app.port}", ""):
            status, _, data = self.request("GET", f"/p/{self.project}", host=host)
            self.assertEqual(status, 403, host)
            self.assertNotIn(b"Cermaq", data)
            status, _, _ = self.request("POST", "/api/projects", {"name": "Intruso"}, JSON, host=host)
            self.assertEqual(status, 403, host)
        self.assertNothingCreated()

    def test_a_request_from_a_page_of_another_site_is_refused(self):
        for origin in ("http://evil.example", "null", f"http://127.0.0.1:{self.app.port + 1}",
                       f"https://127.0.0.1:{self.app.port}"):
            status, _, data = self.request("GET", f"/p/{self.project}", headers={"Origin": origin})
            self.assertEqual(status, 403, origin)
            self.assertNotIn(b"Cermaq", data)
            status, _, _ = self.request("POST", "/api/projects", {"name": "Intruso"}, dict(JSON, Origin=origin))
            self.assertEqual(status, 403, origin)
        for site in ("cross-site", "same-site"):
            status, _, data = self.request("GET", f"/p/{self.project}", headers={"Sec-Fetch-Site": site})
            self.assertEqual(status, 403, site)
            self.assertNotIn(b"Cermaq", data)
            status, _, _ = self.request("GET", f"/p/{self.project}/m/x/f/{FRAME}", headers={"Sec-Fetch-Site": site})
            self.assertEqual(status, 403, site)
            status, _, _ = self.request("POST", "/api/projects", {"name": "Intruso"},
                                        dict(JSON, **{"Sec-Fetch-Site": site}))
            self.assertEqual(status, 403, site)
        self.assertNothingCreated()
        # The page's own requests pass.
        own = {"Origin": self.app.url, "Sec-Fetch-Site": "same-origin"}
        self.assertEqual(self.request("GET", f"/p/{self.project}", headers=own)[0], 200)
        self.assertEqual(self.request("GET", "/", headers={"Sec-Fetch-Site": "none"})[0], 200)

    def test_a_change_sent_as_a_form_is_refused(self):
        form = "name=Intruso&client=x"
        status, _, _ = self.request("POST", "/api/projects", form,
                                    {"Content-Type": "application/x-www-form-urlencoded"})
        self.assertEqual(status, 403)
        status, _, _ = self.request("POST", "/api/projects", form,
                                    {"Content-Type": "application/x-www-form-urlencoded", "X-MeetingTool": "1"})
        self.assertEqual(status, 415)
        status, _, _ = self.request("POST", "/api/projects", json.dumps({"name": "Intruso"}),
                                    {"Content-Type": "text/plain"})
        self.assertEqual(status, 403)
        status, _, _ = self.request("POST", "/api/projects", json.dumps({"name": "Intruso"}),
                                    {"Content-Type": "text/plain", "X-MeetingTool": "1"})
        self.assertEqual(status, 415)
        status, _, _ = self.request("PUT", "/api/upload?kind=transcript&name=t.txt", b"[00:00:01] hola",
                                    {"Content-Type": "application/octet-stream"})
        self.assertEqual(status, 403)
        status, _, _ = self.request("POST", "/api/process", {"project": self.project},
                                    {"Content-Type": "application/json"})
        self.assertEqual(status, 403)
        self.assertNothingCreated()
        self.assertFalse((self.data / ".meetingtool-uploads").exists())
        self.assertEqual(self.app.runner.jobs, {})

    def test_no_permission_is_given_to_another_origin(self):
        status, headers, _ = self.request("OPTIONS", "/api/projects", headers={
            "Origin": "http://evil.example", "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "x-meetingtool,content-type"}, cookie=False)
        self.assertNotEqual(status, 200)
        self.assertNotIn("access-control-allow-origin", headers)
        status, headers, _ = self.request("GET", "/", headers={"Origin": "http://evil.example"})
        self.assertNotIn("access-control-allow-origin", headers)

    def test_the_launch_address_sets_the_cookie_only_with_the_token(self):
        for token in ("", "wrong", self.app.token[:-1], self.app.token + "x"):
            status, headers, _ = self.request("GET", f"/open?token={quote(token)}", cookie=False)
            self.assertEqual(status, 403, token)
            self.assertNotIn("set-cookie", headers)
        status, headers, _ = self.request("GET", f"/open?token={self.app.token}", cookie=False,
                                          headers={"Sec-Fetch-Site": "cross-site"})
        self.assertEqual(status, 403)
        self.assertNotIn("set-cookie", headers)
        status, headers, _ = self.request("GET", f"/open?token={self.app.token}", cookie=False,
                                          headers={"Sec-Fetch-Site": "none"})
        self.assertEqual((status, headers["location"]), (303, "/"))
        cookie = headers["set-cookie"]
        self.assertTrue(cookie.startswith(f"{server.COOKIE}={self.app.token};"))
        self.assertIn("HttpOnly", cookie)
        self.assertIn("SameSite=Strict", cookie)

    def test_each_start_has_a_new_token(self):
        other = server.App(self.tmp / "other", read_key=self.keys.read, save_key=self.keys.save,
                           delete_key=self.keys.delete)
        try:
            self.assertNotEqual(other.token, self.app.token)
            self.assertGreaterEqual(len(other.token), 40)
        finally:
            other.httpd.server_close()

    def test_every_answer_forbids_framing_and_inline_script(self):
        for path, cookie in (("/", True), (f"/p/{self.project}", True), ("/settings", True), ("/nope", True),
                             ("/", False), ("/static/app.js", True)):
            _, headers, data = self.request("GET", path, cookie=cookie)
            self.assertIn("frame-ancestors 'none'", headers["content-security-policy"], path)
            self.assertIn("script-src 'self'", headers["content-security-policy"], path)
            self.assertNotIn("unsafe-inline", headers["content-security-policy"], path)
            self.assertEqual(headers["x-frame-options"], "DENY", path)
            if headers["content-type"].startswith("text/html"):
                text = data.decode("utf-8")
                self.assertEqual(re.findall(r"<script(?![^>]*\bsrc=)", text), [], path)
                self.assertEqual(re.findall(r"\son[a-z]+=", text), [], path)


# ── WI15-AC03: processing from start to end ──────────────────────────────────

def no_usage(first, count):
    """A reading answered with no usage: counted at its maximum."""
    answer = answer_for(first, count)
    del answer["usageMetadata"]
    return answer


class Processing(Running):
    transcript_blocks = test_summary.SPANISH

    def setUp(self):
        super().setUp()
        self.project = self.new_project()
        self.transcript = self.tmp / "Reunión de dudas.docx"
        write_teams_docx(self.transcript, self.transcript_blocks)
        self.video = self.tmp / "grabacion.mp4"
        write_video(self.video, [(SLIDES["A"], 4, False), (SLIDES["B"], 4, False), (SLIDES["C"], 4, False)])

    def process(self, with_recording=True, expect=200, **fields):
        data = {"project": self.project, "title": "Sesión de dudas", "date": "2026-09-25",
                "meeting_type": "requirements", "language": "", "format": "summary", "max_cost": 0.5,
                "transcript": self.upload("transcript", self.transcript)}
        if with_recording:
            data["recording"] = self.upload("recording", self.video)
        data.update(fields)
        answer = self.api("/api/process", data, expect=expect)
        return self.wait(answer["job"]) if expect == 200 else answer

    def wait(self, job_id):
        deadline = time.monotonic() + 180
        while time.monotonic() < deadline:
            status, _, body = self.request("GET", f"/api/jobs/{job_id}")
            self.assertEqual(status, 200)
            job = json.loads(body)
            if job["state"] != "running":
                return job
            time.sleep(0.2)
        self.fail("the job did not end")

    def states(self, job):
        return [(stage["name"], stage["state"]) for stage in job["stages"]]

    def assertNothingLeft(self, job):
        self.assertEqual(job["state"], "failed")
        self.assertEqual(job["meeting"], "")
        self.assertEqual(store.list_meetings(self.data, self.project), [])
        for folder in (library.PROCESSING_DIR, library.RESULTS_DIR):
            leftover = self.data / self.project / folder
            self.assertTrue(not leftover.exists() or not any(leftover.iterdir()), folder)
        self.assertEqual(list((self.data / ".meetingtool-uploads").iterdir()), [])
        self.assertIn("Todavía no hay reuniones", self.page(f"/p/{self.project}"))

    def summary_requests(self, fake=None):
        return [r for r in (fake or self.fake).requests if not r["images"]]


class ProcessTest(Processing):
    def script(self):
        return [lambda first, count: answer_for(first, count), test_summary.returning(test_summary.summary_text(
            "es", "requirements"))]

    def test_the_summary_runs_every_stage_in_order_and_the_meeting_joins_its_project(self):
        job = self.process()
        self.assertEqual(job["state"], "done", job["error"])
        self.assertEqual(self.states(job), [("frames", "done"), ("reading", "done"), ("summary", "done"),
                                            ("report", "done")])
        # The reading's requests come before the summary's, which is the last.
        kinds = ["images" if r["images"] else "text" for r in self.fake.requests]
        self.assertEqual(kinds[-1], "text")
        self.assertNotIn("text", kinds[:-1])
        meetings = store.list_meetings(self.data, self.project)
        self.assertEqual([m["id"] for m in meetings], [job["meeting"]])
        record = meetings[0]
        folder = self.data / self.project / record["folder"]
        self.assertTrue(record["folder"].startswith(f"{library.RESULTS_DIR}/2026-09-25-sesion-de-dudas-"))
        for name in (writer.OUTPUT_NAME, library.REPORT_NAME, library.RUN_NAME, gemini.OUTPUT_NAME,
                     "transcript.docx"):
            self.assertTrue((folder / name).is_file(), name)
        self.assertEqual(record["transcript"], str(folder / "transcript.docx"))
        self.assertGreater(len(list(folder.glob("frame_*.jpg"))), 0)
        self.assertEqual(list(folder.glob("recording*")), [])
        self.assertFalse((self.data / self.project / library.PROCESSING_DIR).exists())
        self.assertEqual(list((self.data / ".meetingtool-uploads").iterdir()), [])
        document.check_report(folder / library.REPORT_NAME, document.summary_headings(
            (folder / writer.OUTPUT_NAME).read_text(encoding="utf-8")), [], [])
        run = json.loads((folder / library.RUN_NAME).read_text(encoding="utf-8"))
        self.assertAlmostEqual(run["cost_usd"], job["spent_usd"], places=4)
        self.assertAlmostEqual(sum(stage["cost_usd"] for stage in job["stages"]), job["spent_usd"], places=3)
        self.assertGreater(job["stages"][1]["cost_usd"], 0)
        self.assertGreater(job["stages"][2]["cost_usd"], 0)
        self.assertEqual(run["requests"], len(self.fake.requests))
        text = self.page(f"/p/{self.project}/m/{job['meeting']}")
        self.assertIn("Resumen ejecutivo", text)
        self.assertIn("Lo que costó", text)
        self.assertIn("Relevamiento", self.page(f"/p/{self.project}"))

    def test_the_summary_request_is_the_one_the_command_sends(self):
        twin = self.tmp / "twin-data"
        store.create_project(twin, "Cermaq Sprint 3", "Cermaq")
        job = self.process()
        self.assertEqual(job["state"], "done", job["error"])
        record = store.list_meetings(self.data, self.project)[0]
        copy = self.tmp / "command-folder"
        shutil.copytree(self.data / self.project / record["folder"], copy)
        for name in (writer.OUTPUT_NAME, library.REPORT_NAME, library.RUN_NAME):
            (copy / name).unlink()
        with FakeGemini([test_summary.returning(test_summary.summary_text("es", "requirements"))]) as command:
            with contextlib.redirect_stdout(io.StringIO()):
                code = summary_main(["--frames", str(copy), "--transcript", str(copy / "transcript.docx"),
                                     "--project", "cermaq-sprint-3", "--title", "Sesión de dudas", "--date",
                                     "2026-09-25", "--type", "requirements", "--data-dir", str(twin)],
                                    read_key=lambda: KEY, endpoint=command.endpoint, sleep=lambda s: None)
            self.assertEqual(code, 0)
            self.assertEqual(self.summary_requests(command)[0]["body"], self.summary_requests()[-1]["body"])
        self.assertEqual(store.list_meetings(twin, "cermaq-sprint-3")[0]["summary"], record["summary"])
        self.assertEqual(store.list_meetings(twin, "cermaq-sprint-3")[0]["key_points"], record["key_points"])

    def test_the_report_cover_gets_the_projects_client_and_the_meetings_type(self):
        """WI16-AC01, from the application."""
        document.set_template(test_report.fields_template(self.tmp / "portada.docx", [
            "{cliente} · {proyecto}", "{reunion} · {tipo} · {fecha}"]), self.data)
        job = self.process()
        self.assertEqual(job["state"], "done", job["error"])
        (record,) = store.list_meetings(self.data, self.project)
        report = docx.Document(str(self.data / self.project / record["folder"] / library.REPORT_NAME))
        self.assertEqual([p.text for p in report.paragraphs[:2]],
                         ["Cermaq · Cermaq Sprint 3", "Sesión de dudas · Relevamiento · 25 de septiembre de 2026"])

    def test_the_summary_needs_the_recording(self):
        answer = self.process(with_recording=False, expect=400)
        self.assertIn("el resumen necesita el video", answer["error"])
        self.assertEqual(self.app.runner.jobs, {})
        self.assertEqual(self.fake.requests, [])

    def test_a_request_that_cannot_be_processed_is_refused_before_anything_is_done(self):
        for fields, said in (({"project": "no-existe"}, "no hay un proyecto"), ({"title": " "}, "título"),
                             ({"date": "25/09/2026"}, "fecha"), ({"date": "2026-02-30"}, "fecha"),
                             ({"meeting_type": "discovery"}, "tipo"), ({"language": "fr"}, "idioma"),
                             ({"format": "word"}, "formato"), ({"max_cost": 0}, "techo"),
                             ({"max_cost": "0.5"}, "techo"), ({"max_cost": 50}, "techo"),
                             ({"transcript": "../../x.docx"}, "transcripción"),
                             ({"recording": "0123456789abcdef01234567.mp4"}, "video")):
            answer = self.process(expect=400, **fields)
            self.assertIn(said, answer["error"], fields)
        self.assertEqual(self.app.runner.jobs, {})
        self.assertEqual(self.fake.requests, [])
        self.assertEqual(store.list_meetings(self.data, self.project), [])

    def test_without_a_key_nothing_is_processed(self):
        self.keys.key = None
        answer = self.process(expect=400)
        self.assertIn("no hay una clave de Gemini", answer["error"])
        self.assertEqual(self.fake.requests, [])

    def test_one_run_at_a_time(self):
        release = threading.Event()
        real = __import__("meetingtool.frames.extract", fromlist=["extract_frames"]).extract_frames

        def slow(*args, **kwargs):
            release.wait(60)
            return real(*args, **kwargs)

        with mock.patch("meetingtool.frames.extract.extract_frames", side_effect=slow):
            first = self.api("/api/process", {
                "project": self.project, "title": "Primera", "date": "2026-09-25", "format": "summary",
                "max_cost": 0.5, "transcript": self.upload("transcript", self.transcript),
                "recording": self.upload("recording", self.video)})["job"]
            self.assertIn("Primera", self.request("GET", "/api/running")[2].decode("utf-8"))
            second = self.process(expect=400, title="Segunda")
            self.assertIn("ya hay una reunión procesándose", second["error"])
            release.set()
            self.assertEqual(self.wait(first)["state"], "done")
        self.assertEqual(json.loads(self.request("GET", "/api/running")[2]), {"job": None})
        self.assertIn(f'data-job="{first}"', self.page(f"/job/{first}"))


class QAProcessTest(Processing):
    transcript_blocks = test_qa.SPANISH

    def script(self):
        return [test_qa.json_answer(test_qa.verbal())]

    def test_the_register_without_a_recording_reads_no_frame(self):
        job = self.process(with_recording=False, format="qa")
        self.assertEqual(job["state"], "done", job["error"])
        self.assertEqual(self.states(job), [("frames", "skipped"), ("reading", "skipped"), ("qa", "done"),
                                            ("report", "done")])
        self.assertEqual([r for r in self.fake.requests if r["images"]], [])
        record = store.list_meetings(self.data, self.project)[0]
        folder = self.data / self.project / record["folder"]
        self.assertTrue((folder / qa.REGISTER_NAME).is_file())
        self.assertTrue((folder / library.REPORT_NAME).is_file())
        self.assertTrue(record["summary"].startswith("Preguntas y respuestas: 2 ("), record["summary"])
        text = self.page(f"/p/{self.project}/m/{record['id']}")
        self.assertIn("Preguntas y respuestas", text)
        self.assertIn("¿Quién carga los estándares de horas por kilo en la planilla?", text)

    def test_the_register_request_is_the_one_the_command_sends(self):
        twin = self.tmp / "twin-data"
        store.create_project(twin, "Cermaq Sprint 3", "Cermaq")
        job = self.process(with_recording=False, format="qa", language="es")
        self.assertEqual(job["state"], "done", job["error"])
        record = store.list_meetings(self.data, self.project)[0]
        copy = self.tmp / "command-folder"
        copy.mkdir()
        shutil.copy(self.data / self.project / record["folder"] / "transcript.docx", copy / "transcript.docx")
        with FakeGemini([test_qa.json_answer(test_qa.verbal())]) as command:
            with contextlib.redirect_stdout(io.StringIO()):
                code = summary_main(["--frames", str(copy), "--transcript", str(copy / "transcript.docx"),
                                     "--project", "cermaq-sprint-3", "--title", "Sesión de dudas", "--date",
                                     "2026-09-25", "--type", "requirements", "--language", "es", "--format", "qa",
                                     "--data-dir", str(twin)],
                                    read_key=lambda: KEY, endpoint=command.endpoint, sleep=lambda s: None)
            self.assertEqual(code, 0)
            self.assertEqual(command.requests[0]["body"], self.fake.requests[0]["body"])
        self.assertEqual(store.list_meetings(twin, "cermaq-sprint-3")[0]["key_points"], record["key_points"])


class FailureTest(Processing):
    def test_a_recording_that_cannot_be_read_stops_at_the_frames(self):
        self.video.write_bytes(b"this is not a video")
        job = self.process()
        self.assertEqual(job["failed_stage_name"], "frames")
        self.assertEqual(job["spent_usd"], 0)
        self.assertEqual(self.fake.requests, [])
        self.assertNothingLeft(job)

    def test_a_reading_gemini_refuses_stops_at_the_reading(self):
        self.fake.script[:] = [400]
        job = self.process()
        self.assertEqual(job["failed_stage_name"], "reading")
        self.assertIn("HTTP 400", job["error"])
        self.assertNothingLeft(job)

    def test_a_summary_refused_twice_stops_at_the_summary_and_says_what_was_spent(self):
        broken = test_summary.returning(test_summary.summary_text("es", drop="Decisiones"))
        self.fake.script[:] = [lambda first, count: answer_for(first, count), broken, broken]
        job = self.process()
        self.assertEqual(job["failed_stage_name"], "summary")
        self.assertEqual(self.states(job)[:3], [("frames", "done"), ("reading", "done"), ("summary", "failed")])
        self.assertGreater(job["spent_usd"], 0)
        self.assertAlmostEqual(job["spent_usd"], sum(s["cost_usd"] for s in job["stages"]), places=3)
        self.assertEqual(len(self.summary_requests()), 2)
        self.assertNothingLeft(job)
        self.assertIn(f'data-job="{job["id"]}"', self.page(f"/job/{job['id']}"))

    def test_a_report_that_cannot_be_built_stops_at_the_report(self):
        self.fake.script[:] = [lambda first, count: answer_for(first, count),
                               test_summary.returning(test_summary.summary_text("es"))]
        with mock.patch("meetingtool.report.document.build_report",
                        side_effect=document.ReportError("the report could not be opened again")):
            job = self.process(meeting_type="")
        self.assertEqual(job["failed_stage_name"], "report")
        self.assertIn("could not be opened again", job["error"])
        self.assertGreater(job["spent_usd"], 0)
        self.assertNothingLeft(job)

    def test_a_meeting_that_cannot_be_recorded_leaves_no_folder(self):
        self.fake.script[:] = [lambda first, count: answer_for(first, count),
                               test_summary.returning(test_summary.summary_text("es"))]
        with mock.patch.object(store, "add_meeting", side_effect=store.ProjectError("the disk is full")):
            job = self.process(meeting_type="")
        self.assertEqual(job["failed_stage_name"], "saving")
        self.assertNothingLeft(job)

    def test_one_ceiling_covers_every_stage(self):
        # A reading answered with no usage is counted at its most (about
        # US$0.19 for three frames); the summary could cost about US$0.10 more.
        # Each fits the ceiling alone, not both: the summary is not sent.
        self.fake.script[:] = [no_usage, test_summary.returning(test_summary.summary_text("es"))]
        ceiling = round(gemini.worst_attempt_cost(3) + 0.05, 2)
        self.assertGreater(ceiling, gemini.token_cost(20000 / writer.CHARS_PER_TOKEN, writer.MAX_OUTPUT_TOKENS))
        job = self.process(meeting_type="", max_cost=ceiling)
        self.assertEqual(job["failed_stage_name"], "summary", job["error"])
        self.assertIn("se frenó antes de mandar el resumen", job["error"])  # the application speaks Spanish (WI17)
        self.assertEqual(self.summary_requests(), [])
        self.assertLessEqual(job["spent_usd"], ceiling)
        self.assertNothingLeft(job)


class ReviewFixesTest(Processing):
    """The independent review of afb0d9f: P2-1, P2-2 and the P3s fixed."""

    def script(self):
        return [lambda first, count: answer_for(first, count), test_summary.returning(test_summary.summary_text("es"))]

    def test_a_save_that_fails_half_way_leaves_no_meeting(self):
        # add_meeting writes the record, then the knowledge (P2-1).
        with mock.patch.object(store, "rebuild_knowledge", side_effect=OSError("knowledge.md is locked")):
            job = self.process(meeting_type="")
        self.assertEqual(job["failed_stage_name"], "saving")
        self.assertEqual(list((self.data / self.project / "meetings").glob("*/meeting.json")), [])
        self.assertNothingLeft(job)
        self.assertNotIn("Sesión de dudas", (self.data / self.project / "knowledge.md").read_text(encoding="utf-8"))

    def test_a_name_taken_by_another_meeting_is_not_removed(self):
        real = jobs.secrets.token_hex
        taken = self.data / self.project / library.RESULTS_DIR / "2026-09-25-sesion-de-dudas-abc123"
        taken.mkdir(parents=True)
        (taken / "summary.md").write_text("otra reunión", encoding="utf-8")
        with mock.patch.object(jobs.secrets, "token_hex", side_effect=lambda n: "abc123" if n == 3 else real(n)):
            job = self.process(meeting_type="")
        self.assertEqual(job["state"], "failed")
        self.assertIn("ya hay una carpeta", job["error"])
        self.assertEqual((taken / "summary.md").read_text(encoding="utf-8"), "otra reunión")
        self.assertFalse((self.data / self.project / library.PROCESSING_DIR).exists())

    def test_a_refused_request_leaves_no_upload(self):
        self.process(expect=400, date="mal")
        self.assertEqual(list((self.data / ".meetingtool-uploads").iterdir()), [])

    def test_the_start_clears_what_a_closed_window_left(self):
        left = self.data / self.project / library.PROCESSING_DIR / "2026-09-25-cortada-abc123"
        left.mkdir(parents=True)
        (left / "transcript.docx").write_bytes(b"client data")
        stranger = self.data / "no-es-proyecto" / library.PROCESSING_DIR
        stranger.mkdir(parents=True)
        self.upload("transcript", self.transcript)
        jobs.clear_leftovers(self.data)
        self.assertFalse((self.data / self.project / library.PROCESSING_DIR).exists())
        self.assertFalse((self.data / ".meetingtool-uploads").exists())
        self.assertTrue(stranger.is_dir())  # not a project's: not touched
        self.assertEqual(len(store.list_projects(self.data)), 1)

    def test_one_application_per_data_folder(self):
        lock = server.DataFolderLock(self.data).acquire()
        try:
            with self.assertRaises(server.DataFolderInUse):
                server.DataFolderLock(self.data).acquire()
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()) as said:
                self.assertEqual(server.serve(self.data, open_browser=False), 2)
            self.assertIn("ya está abierto", said.getvalue())
        finally:
            lock.release()
        server.DataFolderLock(self.data).acquire().release()

    def test_the_ceiling_covers_a_retry_of_the_reading_and_of_the_summary(self):
        # The first real run (INGOL D-186): the first request of 70 frames got no
        # answer, was counted at its worst, and its retry did not fit US$0.50.
        reading, summary = gemini.worst_attempt_cost(70), gemini.token_cost(60000 / writer.CHARS_PER_TOKEN,
                                                                            writer.MAX_OUTPUT_TOKENS)
        self.assertGreater(2 * reading, 0.50)
        self.assertGreaterEqual(jobs.DEFAULT_MAX_COST_USD, 2 * reading + 2 * summary)

    def test_a_stop_by_the_ceiling_says_why_the_attempt_before_failed(self):
        counters = gemini.new_counters()
        with mock.patch.object(gemini, "post_generate", return_value=(None, None, "timed out")):
            with self.assertRaises(gemini.ReadingError) as caught:
                gemini.call_checked("http://127.0.0.1:9/x", KEY, {}, lambda answer: answer, 0.26, "frames 1-70",
                                    (0, 0), lambda seconds: None, counters, 0.50)
        self.assertIn("stopped before sending frames 1-70", str(caught.exception))
        self.assertIn("the attempt before got no answer: timed out", str(caught.exception))
        self.assertEqual(counters["attempts"], 1)

    def test_no_other_process_takes_the_port(self):
        # Measured, not a control of this code: on the owner's Windows the thief is refused even
        # without SO_EXCLUSIVEADDRUSE, which the server sets anyway (review P3-6).
        thief = socket.socket()
        try:
            thief.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            with self.assertRaises(OSError):
                thief.bind((server.HOST, self.app.port))
        finally:
            thief.close()


# ── The command ───────────────────────────────────────────────────────────────

class CommandTest(unittest.TestCase):
    def test_meetingtool_app_serves_its_page_on_this_machine(self):
        with tempfile.TemporaryDirectory() as tmp:
            process = subprocess.Popen([sys.executable, "-m", "meetingtool", "app", "--no-browser", "--data-dir",
                                        str(Path(tmp) / "data")], cwd=REPOSITORY, stdout=subprocess.PIPE,
                                       stderr=subprocess.PIPE, text=True, encoding="utf-8")
            try:
                lines = [process.stdout.readline() for _ in range(2)]
                launch = re.search(r"http://127\.0\.0\.1:(\d+)/open\?token=(\S+)", lines[1])
                self.assertIsNotNone(launch, lines)
                connection = http.client.HTTPConnection("127.0.0.1", int(launch.group(1)), timeout=30)
                connection.request("GET", f"/open?token={launch.group(2)}")
                response = connection.getresponse()
                response.read()
                self.assertEqual(response.status, 303)
                cookie = response.getheader("Set-Cookie").split(";")[0]
                connection.request("GET", "/", headers={"Cookie": cookie})
                response = connection.getresponse()
                self.assertEqual(response.status, 200)
                self.assertIn("Proyectos", response.read().decode("utf-8"))
                connection.close()
            finally:
                process.kill()
                process.communicate()

    def test_the_help_names_the_application(self):
        result = subprocess.run([sys.executable, "-m", "meetingtool", "--help"], cwd=REPOSITORY, capture_output=True,
                                text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("app", result.stdout)

    def test_the_page_files_are_packaged(self):
        text = (REPOSITORY / "pyproject.toml").read_text(encoding="utf-8")
        self.assertIn('"meetingtool.app" = ["static/*"]', text)
        for name in server.STATIC_TYPES:
            self.assertTrue((server.STATIC / name).is_file(), name)


if __name__ == "__main__":
    unittest.main()
